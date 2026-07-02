import unittest
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from src.api.users import _user_export_rows, _user_xlsx_response, generate_user_pdf_report
from src.models.database import Base
from src.models.models import Bet, User, user_bets


async def _streaming_response_bytes(response) -> bytes:
    chunks = []
    async for chunk in response.body_iterator:
        chunks.append(chunk if isinstance(chunk, bytes) else chunk.encode("utf-8"))
    return b"".join(chunks)


class UserBetsExportTests(unittest.IsolatedAsyncioTestCase):
    async def test_personal_export_contains_month_day_flat_stake_and_full_bet_columns(self):
        rows = _user_export_rows([
            {
                "resolved_at": "2026-06-10T15:00:00+03:00",
                "event_name": "Team A - Team B",
                "status": "win",
                "coefficient": 1.9,
                "profit_units": 0.9,
                "source_type": "feed",
                "sport_type": "Футбол",
                "bookmaker_names": ["Фонбет"],
                "outcome": "П1",
            }
        ])

        self.assertEqual(rows[0]["month"], "Июнь 2026")
        self.assertEqual(rows[0]["day"], "10.06.2026")
        self.assertEqual(rows[0]["flat_stake"], 1)
        self.assertEqual(rows[0]["result"], "Победа")

        response = _user_xlsx_response(rows, "my-bets.xlsx")
        content = await _streaming_response_bytes(response)
        wb = load_workbook(BytesIO(content))
        ws = wb["My Bets"]
        headers = [cell.value for cell in ws[1]]

        for header in [
            "Месяц",
            "День",
            "Вид спорта",
            "Матч",
            "БК",
            "КФ",
            "Ставка, флет",
            "Прибыль, флеты",
            "Результат",
        ]:
            self.assertIn(header, headers)

        self.assertLess(headers.index("Вид спорта"), headers.index("Матч"))
        self.assertLess(headers.index("БК"), headers.index("КФ"))
        self.assertLess(headers.index("Ставка, флет"), headers.index("Прибыль, флеты"))
        self.assertEqual(ws.cell(row=2, column=headers.index("Ставка, флет") + 1).value, 1)
        self.assertEqual(ws.cell(row=2, column=headers.index("Результат") + 1).value, "Победа")

    async def test_pdf_report_escapes_markup_like_user_and_event_text(self):
        engine = create_async_engine("sqlite+aiosqlite:///:memory:", future=True)
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        Session = async_sessionmaker(engine, expire_on_commit=False)
        try:
            now = datetime.now(timezone.utc)
            bet_id = uuid.uuid4()
            async with Session() as session:
                user = User(
                    telegram_id=12345,
                    first_name="<broken",
                    last_name="Client & Co",
                    username="client>",
                    role="user",
                    stats_display_mode="percent",
                    tg_chat_joined=False,
                    created_at=now,
                    updated_at=now,
                )
                bet = Bet(
                    id=bet_id,
                    event_name="Team <A & Team B>",
                    coefficient=Decimal("1.90"),
                    status="win",
                    created_at=now,
                    resolved_at=now,
                )
                session.add_all([user, bet])
                await session.flush()
                await session.execute(user_bets.insert().values(
                    user_id=user.telegram_id,
                    bet_id=bet.id,
                    access_type="paid_match",
                    match_charged=True,
                ))
                await session.commit()

                response = await generate_user_pdf_report(current_user=user, db=session)
                content = await _streaming_response_bytes(response)

            self.assertEqual(response.media_type, "application/pdf")
            self.assertTrue(content.startswith(b"%PDF"))
        finally:
            async with engine.begin() as conn:
                await conn.run_sync(Base.metadata.drop_all)
            await engine.dispose()


if __name__ == "__main__":
    unittest.main()
