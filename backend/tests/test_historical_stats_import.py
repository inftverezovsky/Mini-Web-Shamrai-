import tempfile
import unittest
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from src.api.admin import get_admin_shamrai_timeline_stats
from src.models.database import Base
from src.models.models import Bet, HistoricalStatsImportBatch, HistoricalStatsMonthly
from src.services.historical_stats import (
    DEFAULT_HISTORICAL_STATS_CUTOFF,
    DEFAULT_HISTORICAL_STATS_SOURCE,
    apply_historical_stats_import,
    load_active_historical_stats_snapshot,
    parse_historical_stats_workbook,
)
from src.services.stats_export import (
    build_stats_export_workbook,
    export_item_from_bet,
    load_shamrai_export_items,
)


def _sample_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "История_до_июня"
    ws.append(["История до июня"])
    ws.append(["Период", "Ставок", "Побед", "Поражений", "Возвратов", "Проходимость", "Оборот", "Прибыль", "ROI", "Ср. кэф", "Топ вид", "Топ БК"])
    ws.append(["Сентябрь", 2, 1, 1, 0, 0.5, 20000, 5000, 0.25, 2.0, "Футбол", "Фонбет"])
    ws.append([])
    ws.append(["БК база до июня — все месяцы"])
    ws.append(["Иконка", "Название", "Ставок", "Побед", "Поражений", "Возвратов", "Оборот", "Прибыль", "Проходимость", "ROI"])
    ws.append(["FB", "Фонбет", 2, 1, 1, 0, 20000, 5000, 0.5, 0.25])
    ws.append([])
    ws.append(["Виды спорта база до июня — все месяцы"])
    ws.append(["Иконка", "Название", "Ставок", "Побед", "Поражений", "Возвратов", "Оборот", "Прибыль", "Проходимость", "ROI"])
    ws.append(["⚽", "Футбол", 2, 1, 1, 0, 20000, 5000, 0.5, 0.25])

    data = wb.create_sheet("Data")
    data.append(["Период", "№", "Вид", "Вид спорта", "Матч", "Иконка БК", "БК", "Коэф.", "Ставка", "Исход", "Оборот", "Прибыль", "Кол-во чел", "Файл-источник"])
    data.append(["Июнь 2026", 1, "🎾", "Теннис", "A - B", "WL", "Винлайн", 1.9, "П1", "Победа", 10000, 9000, None, "sample.xlsx"])
    data.append(["Июнь 2026", 2, "⚽", "Футбол", "C - D", "LS", "ЛигаСтавок", 1.7, "П2", "Поражение", 10000, -10000, None, "sample.xlsx"])

    wb.create_sheet("Общее")
    wb.create_sheet("Срезы")
    wb.save(path)


class HistoricalStatsImportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:", future=True)
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        self.Session = async_sessionmaker(self.engine, expire_on_commit=False)

    async def asyncTearDown(self):
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
        await self.engine.dispose()

    def test_parse_sample_workbook_adds_june_detail_to_historical_baseline(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "stats.xlsx"
            _sample_workbook(source)

            parsed = parse_historical_stats_workbook(source)

        self.assertEqual(parsed.total_bets, 4)
        self.assertEqual(parsed.total_wins, 2)
        self.assertEqual(parsed.total_losses, 2)
        self.assertEqual(parsed.total_refunds, 0)
        self.assertEqual(parsed.total_profit_rub, Decimal("4000.00"))
        self.assertEqual(len(parsed.monthly), 2)
        self.assertEqual(len(parsed.details), 2)
        self.assertEqual([row.period_key for row in parsed.monthly], ["2025-09", "2026-06"])

    def test_parse_default_source_totals_when_local_file_is_available(self):
        if not DEFAULT_HISTORICAL_STATS_SOURCE.exists():
            self.skipTest("Local Shamrai historical workbook is not available")

        parsed = parse_historical_stats_workbook(DEFAULT_HISTORICAL_STATS_SOURCE)

        self.assertEqual(parsed.total_bets, 1397)
        self.assertEqual(parsed.total_wins, 944)
        self.assertEqual(parsed.total_losses, 449)
        self.assertEqual(parsed.total_refunds, 4)
        self.assertEqual(parsed.total_profit_rub, Decimal("3399200.00"))
        self.assertEqual(len(parsed.details), 98)

    async def test_apply_import_is_idempotent_and_dry_run_does_not_mutate(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "stats.xlsx"
            _sample_workbook(source)
            parsed = parse_historical_stats_workbook(source)

            async with self.Session() as session:
                dry_run = await apply_historical_stats_import(session, parsed, apply=False)
                batch_count = await session.scalar(select(func.count(HistoricalStatsImportBatch.id)))
                self.assertEqual(dry_run.status, "dry_run")
                self.assertEqual(batch_count, 0)

                first = await apply_historical_stats_import(session, parsed, apply=True)
                await session.commit()
                second = await apply_historical_stats_import(session, parsed, apply=True)
                await session.commit()

                batch_count = await session.scalar(select(func.count(HistoricalStatsImportBatch.id)))
                month_count = await session.scalar(select(func.count(HistoricalStatsMonthly.id)))
                snapshot = await load_active_historical_stats_snapshot(session, "all")

        self.assertEqual(first.status, "applied")
        self.assertEqual(second.status, "applied")
        self.assertEqual(batch_count, 1)
        self.assertEqual(month_count, 2)
        self.assertIsNotNone(snapshot)
        assert snapshot is not None
        self.assertEqual(snapshot.summary["bets"], 4)

    async def test_shamrai_export_filters_live_rows_before_historical_cutoff(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "stats.xlsx"
            _sample_workbook(source)
            parsed = parse_historical_stats_workbook(source)

            async with self.Session() as session:
                await apply_historical_stats_import(session, parsed, apply=True)
                session.add_all([
                    Bet(
                        id=uuid.uuid4(),
                        event_name="Old live duplicate",
                        coefficient=Decimal("2.00"),
                        status="win",
                        publication_type="forecast",
                        created_at=datetime(2026, 6, 10, tzinfo=timezone.utc),
                        resolved_at=datetime(2026, 6, 10, tzinfo=timezone.utc),
                    ),
                    Bet(
                        id=uuid.uuid4(),
                        event_name="Fresh July live",
                        coefficient=Decimal("2.50"),
                        status="win",
                        publication_type="forecast",
                        created_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
                        resolved_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
                    ),
                ])
                await session.commit()

                items = await load_shamrai_export_items(session, "all")
                snapshot = await load_active_historical_stats_snapshot(session, "all")

        self.assertEqual([item.event_name for item in items], ["Fresh July live"])
        self.assertIsNotNone(snapshot)
        assert snapshot is not None
        self.assertEqual(snapshot.cutoff_at.isoformat(), DEFAULT_HISTORICAL_STATS_CUTOFF.isoformat())

    async def test_workbook_uses_historical_summary_without_fake_detail_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "stats.xlsx"
            _sample_workbook(source)
            parsed = parse_historical_stats_workbook(source)

            async with self.Session() as session:
                await apply_historical_stats_import(session, parsed, apply=True)
                snapshot = await load_active_historical_stats_snapshot(session, "all")

        live_item = export_item_from_bet(Bet(
            id=uuid.uuid4(),
            event_name="Fresh July live",
            coefficient=Decimal("2.50"),
            status="win",
            publication_type="forecast",
            created_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
            resolved_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
            sport_type="Футбол",
            outcome="П1",
        ))
        assert live_item is not None
        content = build_stats_export_workbook(
            [live_item],
            title="СТАТИСТИКА SHAMRAI",
            period_label="Весь период",
            include_client=False,
            historical=snapshot,
        )
        wb = load_workbook(BytesIO(content))

        self.assertEqual(wb["Статистика"]["A5"].value, 5)
        detail_values = [cell.value for row in wb["Детально"].iter_rows(values_only=False) for cell in row if cell.value]
        history_values = [cell.value for row in wb["История"].iter_rows(values_only=False) for cell in row if cell.value]
        self.assertIn("Fresh July live", detail_values)
        self.assertNotIn("Сентябрь 2025", detail_values)
        self.assertIn("Сентябрь 2025", history_values)
        self.assertIn("A - B", history_values)

    async def test_admin_shamrai_timeline_merges_historical_baseline(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "stats.xlsx"
            _sample_workbook(source)
            parsed = parse_historical_stats_workbook(source)

            async with self.Session() as session:
                await apply_historical_stats_import(session, parsed, apply=True)
                session.add_all([
                    Bet(
                        id=uuid.uuid4(),
                        event_name="Old live duplicate",
                        coefficient=Decimal("2.00"),
                        status="win",
                        publication_type="forecast",
                        created_at=datetime(2026, 6, 10, tzinfo=timezone.utc),
                        resolved_at=datetime(2026, 6, 10, tzinfo=timezone.utc),
                    ),
                    Bet(
                        id=uuid.uuid4(),
                        event_name="Fresh July live",
                        coefficient=Decimal("2.50"),
                        status="win",
                        publication_type="forecast",
                        created_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
                        resolved_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
                        sport_type="Футбол",
                        outcome="П1",
                    ),
                ])
                await session.commit()

                payload = await get_admin_shamrai_timeline_stats(period="all", admin=object(), db=session)

        self.assertEqual(payload["summary"]["bets"], 5)
        self.assertEqual(payload["summary"]["wins"], 3)
        self.assertEqual(payload["summary"]["losses"], 2)
        self.assertEqual(payload["summary"]["profit_units"], 1.9)
        self.assertEqual([month["key"] for month in payload["timeline"]], ["2026-07", "2026-06", "2025-09"])
        self.assertEqual(payload["timeline"][1]["days"], [])
        july_bets = payload["timeline"][0]["days"][0]["bets"]
        self.assertEqual([bet["event_name"] for bet in july_bets], ["Fresh July live"])


if __name__ == "__main__":
    unittest.main()
