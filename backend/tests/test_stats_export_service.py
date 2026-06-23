import unittest
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from src.models.database import Base
from src.models.models import Bet, Bookmaker, User, user_bets
from src.services.stats_export import (
    CLIENT_EXPORT_STAKE,
    FLAT_FORMAT,
    XLSX_CALM_PALETTE,
    build_bookmaker_logo_png,
    build_client_info_export_workbook,
    build_stats_export_workbook,
    build_stats_export_workbook_artifact,
    export_item_from_bet,
    load_client_recent_bet_export_rows,
    load_client_info_export_rows,
    load_client_export_groups,
    summarize_export_items,
)


LEGACY_BRIGHT_FILLS = {
    "ECFDF5",
    "FEF2F2",
    "FFFBEB",
    "DCFCE7",
    "DBEAFE",
    "E0F2FE",
    "FFEDD5",
    "FEE2E2",
}


def _bet(*, status, coefficient="2.00", odds_dropped_to=None):
    bookmaker = SimpleNamespace(id=1, name="Фонбет", code="fonbet")
    return SimpleNamespace(
        id=uuid4(),
        event_name="Team A - Team B",
        status=status,
        coefficient=Decimal(coefficient),
        created_at=datetime(2026, 6, 1, 12, tzinfo=timezone.utc),
        resolved_at=datetime(2026, 6, 10, 12, tzinfo=timezone.utc),
        delivery_mode="feed",
        sport_type="Футбол",
        outcome="П1",
        odds_dropped_to=Decimal(str(odds_dropped_to)) if odds_dropped_to is not None else None,
        bookmakers=[bookmaker],
        bookmaker=bookmaker,
    )


def _fill_rgb(cell) -> str:
    value = cell.fill.fgColor.rgb or cell.fill.fgColor.indexed or ""
    return str(value)[-6:].upper()


class StatsExportServiceTests(unittest.TestCase):
    def test_export_summary_uses_rub_unit_stake_and_refunds(self):
        items = [
            export_item_from_bet(_bet(status="win", coefficient="2.50")),
            export_item_from_bet(_bet(status="loss", coefficient="1.80")),
            export_item_from_bet(_bet(status="refund", coefficient="2.00")),
        ]
        summary = summarize_export_items([item for item in items if item])

        self.assertEqual(summary["bets"], 3)
        self.assertEqual(summary["wins"], 1)
        self.assertEqual(summary["losses"], 1)
        self.assertEqual(summary["refunds"], 1)
        self.assertEqual(summary["turnover"], Decimal("20000"))
        self.assertEqual(summary["profit"], Decimal("5000.00"))
        self.assertEqual(summary["winrate"], 0.5)
        self.assertEqual(summary["roi"], Decimal("0.25"))

    def test_shamrai_workbook_uses_rub_stake_and_contains_odds_drop_column(self):
        items = [export_item_from_bet(_bet(status="win", coefficient="1.90", odds_dropped_to="1.55"))]
        content = build_stats_export_workbook(
            [item for item in items if item],
            title="СТАТИСТИКА SHAMRAI",
            period_label="Весь период",
            include_client=False,
        )

        wb = load_workbook(BytesIO(content))
        self.assertEqual(wb.sheetnames, ["Статистика", "Свод по БК и спорту", "Детально"])
        detail = wb["Детально"]
        headers = [cell.value for cell in detail[1]]
        self.assertIn("Упал до", headers)
        self.assertIn("Оборот, ₽", headers)
        self.assertIn("Прибыль, ₽", headers)
        self.assertNotIn("Дата расчета", headers)
        self.assertIn("День", headers)
        self.assertIn("Ставка, флет", headers)
        self.assertIn("Прибыль, флеты", headers)
        self.assertIn("Результат", headers)
        self.assertLess(headers.index("Вид спорта"), headers.index("Матч"))
        self.assertLess(headers.index("БК"), headers.index("Коэфф."))
        self.assertLess(headers.index("Ставка, флет"), headers.index("Результат"))
        day_col = headers.index("День") + 1
        flat_stake_col = headers.index("Ставка, флет") + 1
        flat_profit_col = headers.index("Прибыль, флеты") + 1
        odds_col = headers.index("Упал до") + 1
        profit_col = headers.index("Прибыль, ₽") + 1
        result_col = headers.index("Результат") + 1
        data_row = next(row for row in range(2, detail.max_row + 1) if detail.cell(row=row, column=day_col).value == "10.06.2026")
        self.assertEqual(detail.cell(row=data_row, column=flat_stake_col).value, 1)
        self.assertEqual(detail.cell(row=data_row, column=flat_profit_col).value, 0.9)
        self.assertEqual(detail.cell(row=data_row, column=odds_col).value, 1.55)
        self.assertEqual(detail.cell(row=data_row, column=profit_col).value, 9000)
        self.assertEqual(detail.cell(row=data_row, column=result_col).value, "Победа")
        self.assertIn("2026", [detail.cell(row=row, column=1).value for row in range(2, data_row)])
        self.assertEqual(detail.row_dimensions[data_row].outlineLevel, 4)
        self.assertGreaterEqual(len(detail._images), 1)
        self.assertGreaterEqual(len(wb["Статистика"]._images), 1)
        self.assertGreaterEqual(len(wb["Свод по БК и спорту"]._images), 1)

    def test_google_sheet_workbook_uses_icon_metadata_without_floating_images(self):
        item = export_item_from_bet(_bet(status="win", coefficient="1.90"))
        item.bookmaker_logo_codes = ["fonbet", "fonbet", "winline", "pari", "betboom"]

        artifact = build_stats_export_workbook_artifact(
            [item],
            title="СТАТИСТИКА SHAMRAI",
            period_label="Весь период",
            include_client=False,
            logo_mode="google_cell",
        )

        wb = load_workbook(BytesIO(artifact.xlsx))
        self.assertTrue(all(len(sheet._images) == 0 for sheet in wb.worksheets))
        self.assertGreaterEqual(len(artifact.icon_cells), 3)
        detail_icon = next(
            cell
            for cell in artifact.icon_cells
            if cell.sheet == "Детально" and cell.codes == ("fonbet", "winline", "pari")
        )
        self.assertEqual(detail_icon.column, 7)
        self.assertGreater(detail_icon.row, 4)
        self.assertGreater(detail_icon.width, 0)
        self.assertGreater(detail_icon.height, 0)
        self.assertFalse(any(cell.sheet == "Детально" and cell.row in {2, 3, 4} for cell in artifact.icon_cells))

    def test_bookmaker_logo_png_whitelists_codes_and_limits_count(self):
        logo = build_bookmaker_logo_png(["fonbet", "unknown", "winline", "pari", "betboom"])

        self.assertIsNotNone(logo)
        self.assertEqual(logo.codes, ("fonbet", "winline", "pari"))
        self.assertGreater(len(logo.data), 0)
        self.assertGreater(logo.width, 0)
        self.assertGreater(logo.height, 0)

    def test_stats_workbook_is_drive_ready_and_concise(self):
        june_bet = _bet(status="win", coefficient="2.10")
        july_bet = _bet(status="loss", coefficient="1.80")
        july_bookmaker = SimpleNamespace(id=2, name="Винлайн", code="winline")
        july_bet.event_name = "Team C - Team D"
        july_bet.resolved_at = datetime(2026, 7, 3, 16, tzinfo=timezone.utc)
        july_bet.sport_type = "Теннис"
        july_bet.bookmaker = july_bookmaker
        july_bet.bookmakers = [july_bookmaker]
        items = [export_item_from_bet(bet) for bet in [june_bet, july_bet]]

        content = build_stats_export_workbook(
            [item for item in items if item],
            title="СТАТИСТИКА SHAMRAI",
            period_label="Весь период",
            include_client=False,
        )

        wb = load_workbook(BytesIO(content))
        summary = wb["Статистика"]
        self.assertFalse(summary.sheet_view.showGridLines)
        self.assertIsNotNone(summary.sheet_properties.tabColor)
        self.assertEqual(summary.freeze_panes, "A13")
        self.assertEqual(summary["A1"].value, "СТАТИСТИКА SHAMRAI")
        self.assertEqual(summary["A2"].value, "Период")
        self.assertEqual(summary["B2"].value, "Весь период")
        self.assertEqual(summary["D2"].value, "Дата выгрузки")
        self.assertEqual(summary["G2"].value, "Ед. ставки")
        self.assertIn("₽", str(summary["H2"].value))
        self.assertEqual(summary["A11"].value, "Помесячная сводка")
        self.assertEqual(
            [summary.cell(row=12, column=col).value for col in range(1, 4)],
            ["Месяц", "Ставки", "Побед"],
        )
        self.assertGreaterEqual(float(summary.row_dimensions[1].height), 30)
        self.assertGreaterEqual(float(summary.column_dimensions["A"].width), 14)

        breakdown = wb["Свод по БК и спорту"]
        self.assertEqual(breakdown.freeze_panes, "A3")
        self.assertEqual(
            breakdown.auto_filter.ref,
            f"A2:{get_column_letter(breakdown.max_column)}{breakdown.max_row}",
        )

        detail = wb["Детально"]
        self.assertEqual(detail.freeze_panes, "A2")
        self.assertEqual(detail.auto_filter.ref, f"A1:{get_column_letter(detail.max_column)}{detail.max_row}")
        self.assertGreaterEqual(float(detail.row_dimensions[1].height), 24)
        self.assertGreaterEqual(float(detail.column_dimensions["F"].width), 38)

    def test_stats_workbook_uses_shared_calm_palette(self):
        items = [
            export_item_from_bet(_bet(status="win", coefficient="2.10")),
            export_item_from_bet(_bet(status="loss", coefficient="1.80")),
        ]
        content = build_stats_export_workbook(
            [item for item in items if item],
            title="СТАТИСТИКА SHAMRAI",
            period_label="Весь период",
            include_client=False,
        )

        wb = load_workbook(BytesIO(content))
        palette_values = {color.upper() for color in XLSX_CALM_PALETTE.values()}
        summary = wb["Статистика"]
        detail = wb["Детально"]
        detail_headers = [cell.value for cell in detail[1]]
        day_col = detail_headers.index("День") + 1
        data_row = next(row for row in range(2, detail.max_row + 1) if detail.cell(row=row, column=day_col).value == "10.06.2026")

        for cell in [summary["A1"], summary["A11"], summary["A12"], summary["A4"], summary["A5"], detail["A1"], detail.cell(row=data_row, column=day_col)]:
            self.assertIn(_fill_rgb(cell), palette_values | {"FFFFFF"})
            self.assertNotIn(_fill_rgb(cell), LEGACY_BRIGHT_FILLS)

    def test_client_workbook_uses_flat_labels_and_values(self):
        item = export_item_from_bet(
            _bet(status="win", coefficient="1.90"),
            unit_stake=CLIENT_EXPORT_STAKE,
            client_name="Client",
        )
        content = build_stats_export_workbook(
            [item],
            title="СТАТИСТИКА КЛИЕНТОВ SHAMRAI",
            period_label="Весь период",
            include_client=True,
            value_format=FLAT_FORMAT,
            value_label="флеты",
        )

        wb = load_workbook(BytesIO(content))
        detail = wb["Детально"]
        headers = [cell.value for cell in detail[1]]
        self.assertIn("Оборот, флеты", headers)
        self.assertIn("Прибыль, флеты", headers)
        self.assertNotIn("Дата расчета", headers)
        self.assertIn("День", headers)
        self.assertIn("Ставка, флет", headers)
        profit_col = headers.index("Прибыль, флеты") + 1
        flat_stake_col = headers.index("Ставка, флет") + 1
        data_row = next(row for row in range(2, detail.max_row + 1) if detail.cell(row=row, column=flat_stake_col).value == 1)
        self.assertEqual(detail.cell(row=data_row, column=profit_col).value, 0.9)
        self.assertEqual(detail.cell(row=data_row, column=flat_stake_col).value, 1)
        self.assertGreaterEqual(len(detail._images), 1)

    def test_fonbet_is_primary_bookmaker_in_exports(self):
        fonbet = SimpleNamespace(id=1, name="Фонбет", code="fonbet")
        betboom = SimpleNamespace(id=2, name="БетБум", code="betboom")
        bet = _bet(status="win", coefficient="1.90")
        bet.bookmakers = [betboom, fonbet]
        bet.bookmaker = betboom

        item = export_item_from_bet(bet)

        self.assertEqual(item.bookmaker_names, ["Фонбет"])
        self.assertEqual(item.bookmaker_logo_codes, ["fonbet"])


class StatsExportClientAccessTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:", future=True)
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        self.Session = async_sessionmaker(self.engine, expire_on_commit=False)

    async def asyncTearDown(self):
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
        await self.engine.dispose()

    def _user(self, telegram_id: int, role: str = "user") -> User:
        return User(
            telegram_id=telegram_id,
            username=f"user_{telegram_id}",
            first_name=f"User {telegram_id}",
            role=role,
        )

    def _bet(self, index: int, *, status: str = "win", coefficient: str = "1.90") -> Bet:
        created_at = datetime(2026, 6, 1, 12, tzinfo=timezone.utc) + timedelta(days=index - 1)
        return Bet(
            id=uuid.uuid4(),
            event_name=f"Team {index} A - Team {index} B",
            coefficient=Decimal(coefficient),
            status=status,
            created_at=created_at,
            resolved_at=created_at + timedelta(hours=3),
            sport_type="Футбол",
            outcome="П1",
        )

    async def _add_access(
        self,
        session,
        *,
        user: User,
        bet: Bet,
        access_type: str,
        match_charged: bool,
        taken_at: datetime | None = None,
    ) -> None:
        await session.execute(
            user_bets.insert().values(
                user_id=user.telegram_id,
                bet_id=bet.id,
                access_type=access_type,
                match_charged=match_charged,
                taken_at=taken_at or func.now(),
            )
        )

    async def test_client_export_groups_include_only_paid_non_staff_access(self):
        async with self.Session() as session:
            paid_user = self._user(101)
            staff_user = self._user(900, role="admin")
            users = [paid_user, staff_user]
            bets = [self._bet(index) for index in range(1, 8)]
            session.add_all(users + bets)
            await session.flush()

            await self._add_access(session, user=paid_user, bet=bets[0], access_type="paid_match", match_charged=True)
            await self._add_access(session, user=paid_user, bet=bets[1], access_type="telegram_stars_single_bet", match_charged=False)
            await self._add_access(session, user=paid_user, bet=bets[2], access_type="free_bet", match_charged=False)
            await self._add_access(session, user=paid_user, bet=bets[3], access_type="guarantee_replacement", match_charged=False)
            await self._add_access(session, user=paid_user, bet=bets[4], access_type="crowd_pool", match_charged=False)
            await self._add_access(session, user=paid_user, bet=bets[5], access_type="admin", match_charged=True)
            await self._add_access(session, user=staff_user, bet=bets[6], access_type="paid_match", match_charged=True)
            await session.commit()

            groups = await load_client_export_groups(session, "all")

        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0].user_id, paid_user.telegram_id)
        self.assertEqual([item.access_type for item in groups[0].items], ["paid_match", "telegram_stars_single_bet"])

    async def test_client_info_export_rows_include_all_non_staff_clients(self):
        async with self.Session() as session:
            fonbet = Bookmaker(name="Фонбет", code="fonbet", is_active=True)
            winline = Bookmaker(name="Винлайн", code="winline", is_active=True)
            betboom = Bookmaker(name="БетБум", code="betboom", is_active=True)
            liga = Bookmaker(name="Лига Ставок", code="ligastavok", is_active=True)
            marathon = Bookmaker(name="Марафонбет", code="marathon", is_active=True)
            paid_user = self._user(101)
            paid_user.matches_remaining = 4
            paid_user.phone = "+79990000001"
            paid_user.vk_user_id = "vk-101"
            paid_user.other_bookmaker_name = "Custom BK"
            paid_user.client_group = "VIP"
            paid_user.client_tag = "контроль"
            paid_user.ab_group = "B"
            no_stats_user = self._user(202)
            no_stats_user.purchased_bets_balance = 2
            no_stats_user.ab_group = "A"
            no_stats_user.bookmakers = [fonbet, winline, betboom, liga, marathon]
            staff_user = self._user(900, role="admin")
            bets = [
                self._bet(1, status="win", coefficient="1.90"),
                self._bet(2, status="loss", coefficient="1.90"),
                self._bet(3, status="pending", coefficient="1.90"),
            ]
            bets[2].resolved_at = None
            session.add_all([fonbet, winline, betboom, liga, marathon, paid_user, no_stats_user, staff_user, *bets])
            await session.flush()

            await self._add_access(session, user=paid_user, bet=bets[0], access_type="paid_match", match_charged=True)
            await self._add_access(session, user=paid_user, bet=bets[1], access_type="paid_match", match_charged=True)
            await self._add_access(session, user=paid_user, bet=bets[2], access_type="free_bet", match_charged=False)
            await session.commit()

            rows = await load_client_info_export_rows(session, "all")

        rows_by_id = {row.user_id: row for row in rows}
        self.assertEqual(set(rows_by_id), {101, 202})
        self.assertEqual(rows_by_id[101].matches_remaining, 4)
        self.assertEqual(rows_by_id[101].phone, "+79990000001")
        self.assertEqual(rows_by_id[101].vk_user_id, "vk-101")
        self.assertEqual(rows_by_id[101].bookmaker_names, "Custom BK")
        self.assertEqual(rows_by_id[101].client_group, "VIP")
        self.assertEqual(rows_by_id[101].wins, 1)
        self.assertEqual(rows_by_id[101].losses, 1)
        self.assertEqual(rows_by_id[101].settled_bets, 2)
        self.assertEqual(rows_by_id[101].pending_bets, 1)
        self.assertEqual(rows_by_id[101].total_taken_bets, 3)
        self.assertEqual(rows_by_id[101].recent_results, ["loss", "win"])
        self.assertEqual(rows_by_id[101].situation_label, "Рабочая просадка")
        self.assertEqual(rows_by_id[202].matches_remaining, 2)
        self.assertEqual(rows_by_id[202].bets, 0)
        self.assertEqual(rows_by_id[202].situation_label, "Нет расчетов")

        content = build_client_info_export_workbook(rows, period_label="Весь период")
        wb = load_workbook(BytesIO(content))
        self.assertEqual(wb.sheetnames, ["Обзор", "Клиенты", "По месяцам", "По дням", "Последние ставки"])
        overview = wb["Обзор"]
        self.assertFalse(overview.sheet_view.showGridLines)
        self.assertEqual(overview["A1"].value, "CRM-ОТЧЕТ SHAMRAI")
        self.assertEqual(overview["A2"].value, "Период")
        self.assertEqual(overview["B2"].value, "Весь период")
        self.assertEqual(overview["A15"].value, "Ситуации клиентов")
        self.assertGreaterEqual(float(overview.row_dimensions[1].height), 30)
        self.assertGreaterEqual(len(overview._charts), 2)
        self.assertEqual(overview["H15"].value, "Срез по остаткам абона")
        self.assertIn("Кому срочно написать", [overview.cell(row=row, column=1).value for row in range(1, overview.max_row + 1)])

        ws = wb["Клиенты"]
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Матчей осталось", headers)
        self.assertIn("Телефон", headers)
        self.assertIn("VK ID", headers)
        self.assertIn("БК клиента", headers)
        self.assertNotIn("Иконка БК", headers)
        self.assertIn("Сегмент абона", headers)
        self.assertIn("Рекомендованное действие", headers)
        self.assertIn("Канал связи", headers)
        self.assertIn("A/B", headers)
        self.assertIn("Взял матчей всего", headers)
        self.assertIn("Ожидают расчета", headers)
        self.assertIn("Победы", headers)
        self.assertIn("Поражения", headers)
        self.assertIn("Ситуация", headers)
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(len(ws._images), 0)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertIsNotNone(ws.sheet_properties.tabColor)
        self.assertEqual(ws.auto_filter.ref, f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
        id_col = headers.index("ID") + 1
        bookmaker_col = headers.index("БК клиента") + 1
        username_col = headers.index("Username") + 1
        client_col = headers.index("Клиент") + 1
        segment_col = headers.index("Сегмент абона") + 1
        action_col = headers.index("Рекомендованное действие") + 1
        channel_col = headers.index("Канал связи") + 1
        ab_col = headers.index("A/B") + 1
        date_col = headers.index("Дата регистрации") + 1
        status_col = headers.index("Ситуация") + 1
        description_col = headers.index("Описание") + 1
        paid_row = next(row for row in range(2, ws.max_row + 1) if ws.cell(row=row, column=id_col).value == 101)
        no_stats_row = next(row for row in range(2, ws.max_row + 1) if ws.cell(row=row, column=id_col).value == 202)
        self.assertEqual(ws.cell(row=paid_row, column=segment_col).value, "3+ матча")
        self.assertEqual(ws.cell(row=paid_row, column=action_col).value, "Поддержать клиента")
        self.assertEqual(ws.cell(row=paid_row, column=channel_col).value, "VK")
        self.assertEqual(ws.cell(row=paid_row, column=ab_col).value, "B")
        self.assertEqual(ws.cell(row=no_stats_row, column=segment_col).value, "1-2 матча")
        self.assertEqual(ws.cell(row=no_stats_row, column=action_col).value, "Допродать матчи")
        self.assertEqual(ws.cell(row=no_stats_row, column=channel_col).value, "Telegram")
        self.assertEqual(ws.cell(row=no_stats_row, column=ab_col).value, "A")
        self.assertEqual(
            ws.cell(row=no_stats_row, column=bookmaker_col).value,
            "Фонбет, Винлайн, БетБум, Лига Ставок, Марафонбет",
        )
        self.assertNotIn("+", ws.cell(row=no_stats_row, column=bookmaker_col).value)
        self.assertGreaterEqual(float(ws.column_dimensions[get_column_letter(bookmaker_col)].width), 48)
        self.assertGreaterEqual(float(ws.row_dimensions[no_stats_row].height), 44)
        self.assertLessEqual(float(ws.row_dimensions[no_stats_row].height), 58)
        self.assertFalse(bool(ws.cell(row=no_stats_row, column=username_col).alignment.wrap_text))
        for col in [id_col, date_col, segment_col, action_col, channel_col, ab_col, status_col]:
            self.assertEqual(ws.cell(row=paid_row, column=col).alignment.horizontal, "center")
        for col in [client_col, bookmaker_col, description_col]:
            self.assertTrue(bool(ws.cell(row=paid_row, column=col).alignment.wrap_text))
        palette_values = {color.upper() for color in XLSX_CALM_PALETTE.values()}
        for sheet_name in ["Обзор", "Клиенты", "По месяцам", "По дням", "Последние ставки"]:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    fill = _fill_rgb(cell)
                    if fill and fill != "000000":
                        self.assertNotIn(fill, LEGACY_BRIGHT_FILLS)
            self.assertIn(_fill_rgb(sheet["A1"]), palette_values | {"FFFFFF"})

        monthly = wb["По месяцам"]
        daily = wb["По дням"]
        self.assertEqual([cell.value for cell in monthly[1]][:9], [
            "Месяц",
            "Клиентов",
            "Ставки",
            "Победы",
            "Поражения",
            "Возвраты",
            "Проход",
            "ROI",
            "Профит, флеты",
        ])
        self.assertEqual([cell.value for cell in daily[1]][:9], [
            "День",
            "Клиентов",
            "Ставки",
            "Победы",
            "Поражения",
            "Возвраты",
            "Проход",
            "ROI",
            "Профит, флеты",
        ])
        self.assertEqual(monthly.freeze_panes, "A2")
        self.assertEqual(daily.freeze_panes, "A2")

    async def test_client_recent_bet_export_rows_limit_each_client_to_latest_50_taken_bets(self):
        async with self.Session() as session:
            fonbet = Bookmaker(name="Фонбет", code="fonbet", is_active=True)
            user = self._user(101)
            user.phone = "+79990000101"
            user.vk_user_id = "vk-101"
            user.matches_remaining = 7
            staff_user = self._user(900, role="moderator")
            base_taken_at = datetime(2026, 6, 1, 12, tzinfo=timezone.utc)
            bets = []
            for index in range(60):
                event_time = base_taken_at + timedelta(days=index // 2, minutes=index)
                if index == 57:
                    status = "pending"
                    resolved_at = None
                elif index == 58:
                    status = "refund"
                    resolved_at = event_time + timedelta(hours=2)
                elif index % 2 == 0:
                    status = "win"
                    resolved_at = event_time + timedelta(hours=2)
                else:
                    status = "loss"
                    resolved_at = event_time + timedelta(hours=2)
                bet = self._bet(index + 1, status=status)
                bet.resolved_at = resolved_at
                bet.bookmakers = [fonbet]
                bets.append(bet)
            staff_bet = self._bet(99, status="win")
            staff_bet.bookmakers = [fonbet]
            session.add_all([fonbet, user, staff_user, *bets, staff_bet])
            await session.flush()

            for index, bet in enumerate(bets):
                await self._add_access(
                    session,
                    user=user,
                    bet=bet,
                    access_type="free_bet" if index == 59 else "paid_match",
                    match_charged=index % 3 != 0,
                    taken_at=event_time,
                )
            await self._add_access(
                session,
                user=staff_user,
                bet=staff_bet,
                access_type="paid_match",
                match_charged=True,
                taken_at=base_taken_at + timedelta(days=31, minutes=61),
            )
            await session.commit()

            info_rows = await load_client_info_export_rows(session, "all")
            recent_rows = await load_client_recent_bet_export_rows(session, "all", limit_per_client=50)

        self.assertEqual(len(recent_rows), 50)
        self.assertTrue(all(row.user_id == 101 for row in recent_rows))
        self.assertEqual(recent_rows[0].event_name, "Team 60 A - Team 60 B")
        self.assertEqual(recent_rows[-1].event_name, "Team 11 A - Team 11 B")
        self.assertNotIn("Team 10 A - Team 10 B", [row.event_name for row in recent_rows])
        self.assertIn("pending", {row.status for row in recent_rows})
        self.assertIn("refund", {row.status for row in recent_rows})
        self.assertIn("win", {row.status for row in recent_rows})
        self.assertIn("loss", {row.status for row in recent_rows})
        self.assertEqual(recent_rows[0].phone, "+79990000101")
        self.assertEqual(recent_rows[0].vk_user_id, "vk-101")
        self.assertEqual(recent_rows[0].matches_remaining, 7)
        self.assertEqual(recent_rows[0].access_type, "free_bet")

        content = build_client_info_export_workbook(info_rows, recent_rows=recent_rows, period_label="Весь период")
        wb = load_workbook(BytesIO(content))
        self.assertEqual(wb.sheetnames, ["Обзор", "Клиенты", "По месяцам", "По дням", "Последние ставки"])
        monthly = wb["По месяцам"]
        monthly_headers = [cell.value for cell in monthly[1]]
        month_label_col = monthly_headers.index("Месяц") + 1
        month_bets_col = monthly_headers.index("Ставки") + 1
        month_profit_col = monthly_headers.index("Профит, флеты") + 1
        june_row = next(row for row in range(2, monthly.max_row + 1) if monthly.cell(row=row, column=month_label_col).value == "Июнь 2026")
        self.assertEqual(monthly.cell(row=june_row, column=month_bets_col).value, 50)
        self.assertEqual(monthly.cell(row=june_row, column=month_profit_col).number_format, FLAT_FORMAT)

        daily = wb["По дням"]
        daily_headers = [cell.value for cell in daily[1]]
        day_label_col = daily_headers.index("День") + 1
        day_bets_col = daily_headers.index("Ставки") + 1
        self.assertIn("30.06.2026", [daily.cell(row=row, column=day_label_col).value for row in range(2, daily.max_row + 1)])
        daily_total = next(row for row in range(2, daily.max_row + 1) if daily.cell(row=row, column=day_label_col).value == "ИТОГО")
        self.assertEqual(daily.cell(row=daily_total, column=day_bets_col).value, 50)

        history = wb["Последние ставки"]
        headers = [cell.value for cell in history[1]]
        self.assertIn("Дата взятия", headers)
        self.assertIn("Вид", headers)
        self.assertIn("Вид спорта", headers)
        self.assertIn("Иконка БК", headers)
        self.assertIn("Тип доступа", headers)
        self.assertIn("Матч списан", headers)
        self.assertEqual(history.max_row, 51)
        self.assertGreaterEqual(len(history._images), 1)
        self.assertEqual(history.freeze_panes, "A2")
        self.assertIsNotNone(history.sheet_properties.tabColor)
        self.assertEqual(history.auto_filter.ref, f"A1:{get_column_letter(history.max_column)}{history.max_row}")
        sport_icon_col = headers.index("Вид") + 1
        logo_col = headers.index("Иконка БК") + 1
        match_col = headers.index("Матч") + 1
        bookmaker_col = headers.index("БК") + 1
        bet_id_col = headers.index("ID ставки") + 1
        date_col = headers.index("Дата взятия") + 1
        result_col = headers.index("Результат") + 1
        self.assertEqual(history.cell(row=2, column=sport_icon_col).value, "⚽")
        self.assertIn(history.cell(row=2, column=logo_col).value, (None, ""))
        self.assertEqual(history.cell(row=2, column=logo_col).alignment.horizontal, "center")
        for col in [date_col, result_col, sport_icon_col, logo_col]:
            self.assertEqual(history.cell(row=2, column=col).alignment.horizontal, "center")
        for col in [match_col, bookmaker_col, bet_id_col]:
            self.assertTrue(bool(history.cell(row=2, column=col).alignment.wrap_text))
        self.assertGreaterEqual(float(history.row_dimensions[2].height), 34)
        self.assertGreaterEqual(float(history.column_dimensions[get_column_letter(match_col)].width), 42)


if __name__ == "__main__":
    unittest.main()
