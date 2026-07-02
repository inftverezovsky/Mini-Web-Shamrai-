from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Literal, Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from sqlalchemy.orm import selectinload

from src.core.config import settings
from src.core.roles import STAFF_ROLES
from src.models.models import Bet, User, user_bets
from src.services.historical_stats import HistoricalStatsSnapshot, load_active_historical_stats_snapshot
from src.services.statistics import (
    MONTH_LABELS,
    as_moscow_datetime,
    client_situation,
    is_paid_client_access,
    last_result_codes,
    period_start,
    _prefer_fonbet_bookmakers,
    stat_item_from_bet,
    summarize_items,
)
from src.services.stats_export_styles import (
    COEF_FORMAT,
    COLOR_AMBER_FILL,
    COLOR_BORDER,
    COLOR_DANGER_FILL,
    COLOR_GROUP_DAY,
    COLOR_GROUP_MONTH,
    COLOR_GROUP_YEAR,
    COLOR_HEADER,
    COLOR_HEADER_AMBER,
    COLOR_HEADER_ROSE,
    COLOR_HEADER_SAGE,
    COLOR_HEADER_STEEL,
    COLOR_INFO_FILL,
    COLOR_KPI_LABEL,
    COLOR_ROSE_FILL,
    COLOR_SAGE_FILL,
    COLOR_SUCCESS_FILL,
    COLOR_SURFACE,
    COLOR_SURFACE_ALT,
    COLOR_TEXT,
    COLOR_TEXT_MUTED,
    COLOR_TITLE,
    COLOR_TOTAL_FILL,
    COLOR_WARNING_FILL,
    DATE_FORMAT,
    FLAT_FORMAT,
    MONEY_FORMAT,
    PERCENT_FORMAT,
    XLSX_CALM_PALETTE,
    _profit_fill,
    _solid_fill,
    _status_fill,
    _tone_fill,
)

EXPORT_STATUSES = {"win", "loss", "refund"}
CLIENT_EXPORT_STAKE = Decimal("1")
LogoRenderMode = Literal["floating", "google_cell"]

SPORT_ICONS = {
    "Футбол": "⚽",
    "Баскетбол": "🏀",
    "Гандбол": "🤾",
    "Н/Т": "🏓",
    "Футзал": "🥅",
    "Волейбол": "🏐",
    "Киберспорт": "🎮",
    "Теннис": "🎾",
    "Хоккей": "🏒",
    "Бейсбол": "⚾",
    "Ам. футбол": "🏈",
    "Бадминтон": "🏸",
    "Вод. поло": "🤽",
    "Регби": "🏉",
    "Пляж. футб": "🏖",
    "Крикет": "🏏",
    "Бокс": "🥊",
    "Дартс": "🎯",
    "Единоборства": "🥋",
    "Коньки": "⛸",
    "Лыжи/Биатлон": "🎿",
    "Другие": "✨",
}

BOOKMAKER_CODES = {
    "Фонбет": "FB",
    "Винлайн": "WL",
    "Лига Ставок": "LS",
    "БетБум": "BB",
    "Леон": "LN",
    "Олимпбет": "OL",
    "Марафонбет": "MR",
    "Зенит": "ZN",
}

BOOKMAKER_LOGO_DIR = Path(__file__).resolve().parents[1] / "assets" / "bookmakers"
BOOKMAKER_LOGO_FILES = {
    "fonbet": "fonbet.png",
    "betboom": "betboom.png",
    "winline": "winline.png",
    "pari": "pari.png",
    "ligastavok": "ligastavok.png",
    "marathon": "marathon.png",
    "betcity": "betcity.png",
    "melbet": "melbet.png",
    "leon": "leon.png",
    "olimpbet": "olimpbet.png",
    "zenit": "zenit.png",
    "bettery": "bettery.png",
}
BOOKMAKER_LOGO_MAX_PER_CELL = 3
BOOKMAKER_LOGO_MAX_WIDTH = 34
BOOKMAKER_LOGO_MAX_HEIGHT = 18
BOOKMAKER_LOGO_GAP = 4


def _normalise_bookmaker_key(value: str) -> str:
    return "".join(ch for ch in str(value or "").casefold().replace("ё", "е") if ch.isalnum())


BOOKMAKER_LOGO_ALIASES = {
    "fonbet": "fonbet",
    "фонбет": "fonbet",
    "betboom": "betboom",
    "бетбум": "betboom",
    "winline": "winline",
    "винлайн": "winline",
    "pari": "pari",
    "пари": "pari",
    "ligastavok": "ligastavok",
    "лигаставок": "ligastavok",
    "marathon": "marathon",
    "marathonbet": "marathon",
    "марафон": "marathon",
    "марафонбет": "marathon",
    "betcity": "betcity",
    "бетсити": "betcity",
    "melbet": "melbet",
    "мелбет": "melbet",
    "leon": "leon",
    "леон": "leon",
    "olimp": "olimpbet",
    "olimpbet": "olimpbet",
    "олимп": "olimpbet",
    "олимпбет": "olimpbet",
    "zenit": "zenit",
    "зенит": "zenit",
    "bettery": "bettery",
    "беттери": "bettery",
}
BOOKMAKER_LOGO_CODES_BY_KEY = {
    _normalise_bookmaker_key(key): code for key, code in BOOKMAKER_LOGO_ALIASES.items()
}


@dataclass
class StatsExportItem:
    id: str
    resolved_at: datetime
    period_key: str
    period_label: str
    sport_icon: str
    sport_type: str
    event_name: str
    bookmaker_icon: str
    bookmaker_names: list[str]
    coefficient: Decimal
    odds_dropped_to: Optional[Decimal]
    stake_text: str
    status: str
    result_label: str
    turnover: Decimal
    profit: Decimal
    client_name: str = ""
    source_type: str = ""
    access_type: str = ""
    bookmaker_logo_codes: list[str] = field(default_factory=list)


@dataclass
class ClientStatsExportGroup:
    user_id: int
    client_name: str
    username: str
    items: list[StatsExportItem]


@dataclass
class ClientInfoExportRow:
    user_id: int
    client_name: str
    username: str
    phone: str
    vk_user_id: str
    is_web_only: bool
    bookmaker_names: str
    client_group: str
    client_tag: str
    created_at: Optional[datetime]
    matches_remaining: int
    guarantee_active: bool
    total_taken_bets: int
    settled_bets: int
    pending_bets: int
    refund_bets: int
    bets: int
    wins: int
    losses: int
    winrate: float
    roi: float
    profit_units: float
    average_coefficient: float
    current_streak: int
    current_streak_type: Optional[str]
    max_win_streak: int
    max_loss_streak: int
    recent_results: list[str]
    situation_code: str
    situation_label: str
    situation_tone: str
    situation_description: str
    ab_group: str = ""
    bookmaker_logo_codes: list[str] = field(default_factory=list)
    bookmaker_ids: list[int] = field(default_factory=list)


@dataclass
class ClientRecentBetExportRow:
    user_id: int
    client_name: str
    username: str
    phone: str
    vk_user_id: str
    is_web_only: bool
    client_group: str
    client_tag: str
    matches_remaining: int
    guarantee_active: bool
    taken_at: Optional[datetime]
    event_name: str
    sport_type: str
    bookmaker_names: str
    coefficient: Decimal
    outcome: str
    status: str
    result_label: str
    resolved_at: Optional[datetime]
    source_type: str
    access_type: str
    match_charged: bool
    bet_id: str
    ab_group: str = ""
    bookmaker_logo_codes: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class BookmakerLogoPng:
    data: bytes
    width: int
    height: int
    codes: tuple[str, ...]


@dataclass(frozen=True)
class GoogleSheetIconCell:
    sheet: str
    row: int
    column: int
    codes: tuple[str, ...]
    width: int
    height: int


@dataclass
class StatsWorkbookArtifact:
    xlsx: bytes
    icon_cells: list[GoogleSheetIconCell] = field(default_factory=list)


@dataclass
class _LogoRenderContext:
    mode: LogoRenderMode = "floating"
    icon_cells: list[GoogleSheetIconCell] = field(default_factory=list)


def export_unit_stake() -> Decimal:
    return Decimal(str(settings.STATS_EXPORT_UNIT_STAKE_RUB or 10000))


def _round_money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _display_user(user: User) -> str:
    parts = [part for part in [user.first_name, user.last_name] if part]
    if parts:
        return " ".join(parts)
    if user.username:
        return f"@{user.username}"
    return str(user.telegram_id)


def _user_match_balance(user: User) -> int:
    return int(
        user.purchased_bets_balance
        if (user.purchased_bets_balance or 0) != 0
        else (user.matches_remaining or 0)
    )


def _bookmakers_for_user(user: User) -> str:
    return ", ".join(bookmaker["name"] for bookmaker in _bookmaker_refs_for_user(user))


def _bookmaker_refs_for_user(user: User) -> list[dict[str, Any]]:
    bookmakers: list[dict[str, Any]] = []
    seen = set()
    for bookmaker in list(getattr(user, "bookmakers", None) or []):
        name = str(getattr(bookmaker, "name", "") or "").strip()
        code = str(getattr(bookmaker, "code", "") or "").strip()
        key = (code or name).lower()
        if name and key not in seen:
            seen.add(key)
            bookmakers.append({"id": getattr(bookmaker, "id", None), "name": name, "code": code})
    other_name = str(getattr(user, "other_bookmaker_name", "") or "").strip()
    if other_name and other_name.lower() not in seen:
        bookmakers.append({"name": other_name, "code": ""})
    return bookmakers


def _bookmakers_for_bet(bet: Bet) -> list[dict[str, Any]]:
    bookmakers = []
    seen_ids = set()
    for bookmaker in list(getattr(bet, "bookmakers", None) or []):
        if not bookmaker or bookmaker.id in seen_ids:
            continue
        seen_ids.add(bookmaker.id)
        bookmakers.append({"id": bookmaker.id, "name": bookmaker.name, "code": bookmaker.code})

    fallback = getattr(bet, "bookmaker", None)
    if fallback and fallback.id not in seen_ids:
        bookmakers.append({"id": fallback.id, "name": fallback.name, "code": fallback.code})
    return _prefer_fonbet_bookmakers(bookmakers)


def _bookmaker_logo_code(*, name: str = "", code: str = "") -> Optional[str]:
    for value in (code, name):
        key = _normalise_bookmaker_key(value)
        logo_code = BOOKMAKER_LOGO_CODES_BY_KEY.get(key)
        if logo_code and (BOOKMAKER_LOGO_DIR / BOOKMAKER_LOGO_FILES[logo_code]).is_file():
            return logo_code
    return None


def _bookmaker_logo_codes_for_refs(bookmakers: Iterable[dict[str, Any]]) -> list[str]:
    codes: list[str] = []
    seen = set()
    for bookmaker in bookmakers:
        logo_code = _bookmaker_logo_code(
            name=str(bookmaker.get("name", "") or ""),
            code=str(bookmaker.get("code", "") or ""),
        )
        if logo_code and logo_code not in seen:
            seen.add(logo_code)
            codes.append(logo_code)
    return codes


def _bookmaker_logo_codes_for_names(bookmaker_names: Iterable[str]) -> list[str]:
    return _bookmaker_logo_codes_for_refs({"name": name, "code": ""} for name in bookmaker_names)


def _compact_list_label(values: Iterable[str] | str, *, max_items: int = 3) -> str:
    if isinstance(values, str):
        items = [item.strip() for item in values.split(",")]
    else:
        items = [str(item or "").strip() for item in values]
    clean_items = [item for item in items if item]
    if len(clean_items) <= max_items:
        return ", ".join(clean_items)
    return f"{', '.join(clean_items[:max_items])} +{len(clean_items) - max_items}"


def _logo_cell_text(codes: Iterable[str], fallback: str = "") -> str:
    return "" if _unique_bookmaker_logo_codes(codes) else fallback


def _month_label(value: datetime) -> str:
    return f"{MONTH_LABELS.get(value.month, value.strftime('%m'))} {value.year}"


def _status_label(status: str) -> str:
    if status == "win":
        return "Победа"
    if status == "loss":
        return "Неудача"
    if status == "pending":
        return "Ожидает расчета"
    return "Возврат"


def _profit_for_status(status: str, coefficient: Decimal, unit_stake: Decimal) -> Decimal:
    if status == "win":
        return _round_money((coefficient - Decimal("1")) * unit_stake)
    if status == "loss":
        return -unit_stake
    return Decimal("0")


def export_item_from_bet(
    bet: Bet,
    *,
    unit_stake: Optional[Decimal] = None,
    client_name: str = "",
    access_type: str = "",
) -> Optional[StatsExportItem]:
    if str(getattr(bet, "publication_type", "forecast") or "forecast") != "forecast":
        return None
    if bet.status not in EXPORT_STATUSES or not bet.resolved_at:
        return None

    resolved_at = as_moscow_datetime(bet.resolved_at)
    if not resolved_at:
        return None
    coefficient = Decimal(str(bet.coefficient or "0"))
    stake = unit_stake if unit_stake is not None else export_unit_stake()
    turnover = stake if bet.status in {"win", "loss"} else Decimal("0")
    profit = _profit_for_status(str(bet.status), coefficient, stake)
    bookmakers = _bookmakers_for_bet(bet)
    bookmaker_names = [bookmaker["name"] for bookmaker in bookmakers] or ["Без БК"]
    first_bookmaker = bookmaker_names[0]
    sport_type = str(getattr(bet, "sport_type", None) or "Без спорта")
    delivery_mode = str(getattr(bet, "delivery_mode", None) or "feed")
    if delivery_mode == "feed":
        source_type = "feed"
    elif delivery_mode == "paid_set":
        source_type = "paid_set"
    else:
        source_type = "private"
    odds_dropped_to = getattr(bet, "odds_dropped_to", None)

    return StatsExportItem(
        id=str(bet.id),
        resolved_at=resolved_at,
        period_key=resolved_at.strftime("%Y-%m"),
        period_label=_month_label(resolved_at),
        sport_icon=SPORT_ICONS.get(sport_type, "✨"),
        sport_type=sport_type,
        event_name=str(getattr(bet, "event_name", "") or ""),
        bookmaker_icon=BOOKMAKER_CODES.get(first_bookmaker, ""),
        bookmaker_names=bookmaker_names,
        coefficient=coefficient,
        odds_dropped_to=Decimal(str(odds_dropped_to)) if odds_dropped_to is not None else None,
        stake_text=str(getattr(bet, "outcome", None) or ""),
        status=str(bet.status),
        result_label=_status_label(str(bet.status)),
        turnover=turnover,
        profit=profit,
        client_name=client_name,
        source_type=source_type,
        access_type=access_type,
        bookmaker_logo_codes=_bookmaker_logo_codes_for_refs(bookmakers),
    )


def _summary_from_parts(
    *,
    total: int,
    wins: int,
    losses: int,
    refunds: int,
    turnover: Decimal,
    profit: Decimal,
    coefficient_sum: Decimal,
    coefficient_count: int,
) -> dict[str, Any]:
    resolved = wins + losses
    return {
        "bets": total,
        "wins": wins,
        "losses": losses,
        "refunds": refunds,
        "buyouts": 0,
        "winrate": (wins / resolved) if resolved else 0,
        "average_coefficient": (coefficient_sum / Decimal(coefficient_count)) if coefficient_count else Decimal("0"),
        "turnover": turnover,
        "profit": profit,
        "roi": (profit / turnover) if turnover else 0,
    }


def summarize_export_items(
    items: Iterable[StatsExportItem],
    *,
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> dict[str, Any]:
    item_list = list(items)
    total = len(item_list)
    wins = sum(1 for item in item_list if item.status == "win")
    losses = sum(1 for item in item_list if item.status == "loss")
    refunds = sum(1 for item in item_list if item.status == "refund")
    turnover = sum((item.turnover for item in item_list), Decimal("0"))
    profit = sum((item.profit for item in item_list), Decimal("0"))
    coefficient_count = sum(1 for item in item_list if item.coefficient > 0)
    coefficient_sum = sum((item.coefficient for item in item_list), Decimal("0"))

    if historical:
        summary = historical.summary
        total += int(summary.get("bets") or 0)
        wins += int(summary.get("wins") or 0)
        losses += int(summary.get("losses") or 0)
        refunds += int(summary.get("refunds") or 0)
        turnover += Decimal(str(summary.get("turnover_rub") or "0"))
        profit += Decimal(str(summary.get("profit_rub") or "0"))
        historical_count = sum(int(row.get("bets") or 0) for row in historical.monthly)
        historical_sum = sum(
            Decimal(str(row.get("average_coefficient") or "0")) * Decimal(int(row.get("bets") or 0))
            for row in historical.monthly
        )
        coefficient_count += historical_count
        coefficient_sum += historical_sum

    return _summary_from_parts(
        total=total,
        wins=wins,
        losses=losses,
        refunds=refunds,
        turnover=turnover,
        profit=profit,
        coefficient_sum=coefficient_sum,
        coefficient_count=coefficient_count,
    )


def _flat_stake_for_item(item: StatsExportItem) -> Decimal:
    return Decimal("1") if item.status in {"win", "loss"} else Decimal("0")


def _flat_profit_for_item(item: StatsExportItem) -> Decimal:
    if item.status == "win":
        return item.coefficient - Decimal("1")
    if item.status == "loss":
        return Decimal("-1")
    return Decimal("0")


def _sorted_items(items: Iterable[StatsExportItem]) -> list[StatsExportItem]:
    return sorted(items, key=lambda item: (item.resolved_at, item.event_name))


def _parts_from_items(items: Iterable[StatsExportItem]) -> dict[str, Any]:
    item_list = list(items)
    return {
        "bets": len(item_list),
        "wins": sum(1 for item in item_list if item.status == "win"),
        "losses": sum(1 for item in item_list if item.status == "loss"),
        "refunds": sum(1 for item in item_list if item.status == "refund"),
        "turnover": sum((item.turnover for item in item_list), Decimal("0")),
        "profit": sum((item.profit for item in item_list), Decimal("0")),
        "coefficient_sum": sum((item.coefficient for item in item_list if item.coefficient > 0), Decimal("0")),
        "coefficient_count": sum(1 for item in item_list if item.coefficient > 0),
    }


def _merge_summary_parts(parts: Iterable[dict[str, Any]]) -> dict[str, Any]:
    total_parts = list(parts)
    return _summary_from_parts(
        total=sum(int(part.get("bets") or 0) for part in total_parts),
        wins=sum(int(part.get("wins") or 0) for part in total_parts),
        losses=sum(int(part.get("losses") or 0) for part in total_parts),
        refunds=sum(int(part.get("refunds") or 0) for part in total_parts),
        turnover=sum((Decimal(str(part.get("turnover") or "0")) for part in total_parts), Decimal("0")),
        profit=sum((Decimal(str(part.get("profit") or "0")) for part in total_parts), Decimal("0")),
        coefficient_sum=sum((Decimal(str(part.get("coefficient_sum") or "0")) for part in total_parts), Decimal("0")),
        coefficient_count=sum(int(part.get("coefficient_count") or 0) for part in total_parts),
    )


def _historical_summary_part(row: dict[str, Any]) -> dict[str, Any]:
    bets = int(row.get("bets") or 0)
    return {
        "bets": bets,
        "wins": int(row.get("wins") or 0),
        "losses": int(row.get("losses") or 0),
        "refunds": int(row.get("refunds") or 0),
        "turnover": Decimal(str(row.get("turnover") or "0")),
        "profit": Decimal(str(row.get("profit") or "0")),
        "coefficient_sum": Decimal(str(row.get("average_coefficient") or "0")) * Decimal(bets),
        "coefficient_count": bets,
    }


def _monthly_rows(
    items: list[StatsExportItem],
    *,
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> list[dict[str, Any]]:
    months: dict[str, dict[str, Any]] = {}
    for row in (historical.monthly if historical else []):
        key = str(row.get("period_key") or "")
        if not key:
            continue
        months.setdefault(key, {"label": row.get("period_label") or key, "parts": []})["parts"].append(_historical_summary_part(row))

    live_months: dict[str, list[StatsExportItem]] = defaultdict(list)
    for item in items:
        live_months[item.period_key].append(item)
    for key, month_items in live_months.items():
        months.setdefault(key, {"label": month_items[0].period_label, "parts": []})["parts"].append(_parts_from_items(month_items))

    rows = []
    total_parts = []
    for key in sorted(months):
        summary = _merge_summary_parts(months[key]["parts"])
        total_parts.extend(months[key]["parts"])
        rows.append({
            "label": months[key]["label"],
            **summary,
        })
    rows.append({"label": "ИТОГО", **_merge_summary_parts(total_parts)})
    return rows


def _historical_breakdown_source(
    historical: Optional[HistoricalStatsSnapshot],
    *,
    kind: str,
) -> list[dict[str, Any]]:
    if not historical:
        return []
    return historical.bookmaker_breakdowns if kind == "bookmaker" else historical.sport_breakdowns


def _breakdown_rows(
    items: list[StatsExportItem],
    *,
    kind: str,
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    for row in _historical_breakdown_source(historical, kind=kind):
        key = str(row.get("normalized_key") or row.get("label") or "").lower()
        if not key:
            continue
        label = str(row.get("label") or "")
        groups[key] = {
            "label": label,
            "icon": row.get("icon") or (BOOKMAKER_CODES.get(label, "") if kind == "bookmaker" else ""),
            "logo_codes": [row.get("bookmaker_code")] if kind == "bookmaker" and row.get("bookmaker_code") not in {"", "other", None} else [],
            "parts": [_historical_summary_part(row)],
        }

    for item in items:
        labels = item.bookmaker_names if kind == "bookmaker" else [item.sport_type]
        for label in labels:
            key = label.lower()
            if key not in groups:
                groups[key] = {
                    "label": label,
                    "icon": BOOKMAKER_CODES.get(label, "") if kind == "bookmaker" else item.sport_icon,
                    "logo_codes": _bookmaker_logo_codes_for_names([label]) if kind == "bookmaker" else [],
                    "parts": [],
                }
            groups[key]["parts"].append(_parts_from_items([item]))

    rows = [
        {
            "icon": group["icon"],
            "label": group["label"],
            "logo_codes": group["logo_codes"],
            **_merge_summary_parts(group["parts"]),
        }
        for group in groups.values()
    ]
    return sorted(rows, key=lambda row: (-int(row["bets"]), str(row["label"])))


def _safe_sheet_title(title: str) -> str:
    clean = "".join(ch for ch in title if ch not in r'[]:*?/\\').strip() or "Статистика"
    return clean[:31]


def _style_title(ws, max_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    ws["A1"].fill = _solid_fill(COLOR_TITLE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30


def _style_range_header(ws, row: int, start_col: int, end_col: int, fill: str = COLOR_HEADER) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _solid_fill(fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 15), 24)


def _format_stake_unit(value_label: str) -> str:
    if value_label == "₽":
        stake = export_unit_stake()
        return f"{int(stake):,}".replace(",", " ") + " ₽"
    if value_label == "флеты":
        return "1 флет"
    return value_label


def _generated_at_text() -> str:
    generated_at = as_moscow_datetime(datetime.now(timezone.utc))
    return generated_at.strftime("%d.%m.%Y") if generated_at else ""


def _write_metadata_row(ws, row: int, period_label: str, *, value_label: str) -> None:
    metadata = [
        (1, "Период", 2, period_label),
        (4, "Дата выгрузки", 5, _generated_at_text()),
        (7, "Ед. ставки", 8, _format_stake_unit(value_label)),
    ]
    for col in range(1, 14):
        cell = ws.cell(row=row, column=col)
        cell.fill = _solid_fill(COLOR_SURFACE)
        cell.alignment = Alignment(vertical="center")
    for label_col, label, value_col, value in metadata:
        label_cell = ws.cell(row=row, column=label_col, value=label)
        value_cell = ws.cell(row=row, column=value_col, value=value)
        label_cell.font = Font(bold=True, color=COLOR_TEXT_MUTED)
        value_cell.font = Font(bold=True, color=COLOR_TEXT)
        label_cell.alignment = Alignment(horizontal="right", vertical="center")
        value_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22


def _write_section_title(ws, row: int, start_col: int, end_col: int, title: str, *, fill: str) -> None:
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = ws.cell(row=row, column=start_col, value=title)
    cell.font = Font(bold=True, size=12, color=COLOR_TEXT)
    cell.fill = _solid_fill(fill)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22


def _write_export_group_row(ws, row: int, max_col: int, label: str, *, level: int, fill: str) -> None:
    ws.cell(row=row, column=1, value=label)
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color=COLOR_TEXT)
        cell.fill = _solid_fill(fill)
        cell.alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row].outlineLevel = level


def _set_auto_filter(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max(max_row, min_row)}"
    ws.auto_filter.ref = f"{start}:{end}"


def _set_column_widths(ws, widths: dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _set_border(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    side = Side(style="thin", color=COLOR_BORDER)
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _polish_table_header(ws, row: int, start_col: int, end_col: int) -> None:
    ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 15), 30)
    for col in range(start_col, end_col + 1):
        ws.cell(row=row, column=col).alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
            shrink_to_fit=True,
        )


def _polish_table_body(
    ws,
    min_row: int,
    max_row: int,
    max_col: int,
    *,
    center_cols: set[int],
    wrap_cols: set[int] | None = None,
    right_cols: set[int] | None = None,
    row_height: float = 34,
) -> None:
    wrap_cols = wrap_cols or set()
    right_cols = right_cols or set()
    for row in range(min_row, max_row + 1):
        has_wrapped_content = any(str(ws.cell(row=row, column=col).value or "") for col in wrap_cols)
        effective_height = row_height + (8 if has_wrapped_content and wrap_cols else 0)
        ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 15), effective_height)
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            if col in right_cols:
                horizontal = "right"
            elif col in center_cols:
                horizontal = "center"
            else:
                horizontal = "left"
            cell.alignment = Alignment(
                horizontal=horizontal,
                vertical="center",
                wrap_text=col in wrap_cols,
                shrink_to_fit=col not in wrap_cols,
            )


def _truncate_cell_text(value: str, *, limit: int) -> str:
    clean = " ".join(str(value or "").split())
    if len(clean) <= limit:
        return clean
    return clean[: max(0, limit - 1)].rstrip() + "…"


def _apply_date_row_groups(ws, row_dates: list[tuple[int, Optional[datetime]]]) -> None:
    dated_rows = [(row, as_moscow_datetime(value)) for row, value in row_dates if as_moscow_datetime(value)]
    if not dated_rows:
        return

    ws.sheet_properties.outlinePr.summaryBelow = False

    def group_by(key_func, level: int) -> None:
        start_row: Optional[int] = None
        end_row: Optional[int] = None
        previous_key = object()
        for row, value in dated_rows:
            key = key_func(value)
            if start_row is None:
                start_row = end_row = row
                previous_key = key
                continue
            if key == previous_key and end_row is not None and row == end_row + 1:
                end_row = row
                continue
            if start_row is not None and end_row is not None:
                ws.row_dimensions.group(start_row, end_row, outline_level=level, hidden=False)
            start_row = end_row = row
            previous_key = key
        if start_row is not None and end_row is not None:
            ws.row_dimensions.group(start_row, end_row, outline_level=level, hidden=False)

    group_by(lambda value: value.year, 1)
    group_by(lambda value: (value.year, value.month), 2)
    group_by(lambda value: (value.year, value.month, value.day), 3)


def _format_summary_value(cell, key: str, value_format: str = MONEY_FORMAT) -> None:
    if key in {"winrate", "roi"}:
        cell.number_format = PERCENT_FORMAT
    elif key in {"turnover", "profit"}:
        cell.number_format = value_format
    elif key == "average_coefficient":
        cell.number_format = COEF_FORMAT


def _write_kpi(
    ws,
    start_row: int,
    start_col: int,
    label: str,
    value: Any,
    key: str,
    *,
    value_format: str = MONEY_FORMAT,
) -> None:
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 1, start_column=start_col, end_row=start_row + 1, end_column=start_col + 1)
    label_cell = ws.cell(row=start_row, column=start_col, value=label)
    value_cell = ws.cell(row=start_row + 1, column=start_col, value=value)
    label_cell.font = Font(bold=True, color=COLOR_TEXT_MUTED)
    label_cell.fill = _solid_fill(COLOR_KPI_LABEL)
    label_cell.alignment = Alignment(horizontal="center")
    value_cell.font = Font(bold=True, size=14, color=COLOR_TEXT)
    if key in {"profit", "roi"} and value:
        value_cell.fill = _profit_fill(value)
    else:
        value_cell.fill = _solid_fill(COLOR_SURFACE)
    ws.row_dimensions[start_row].height = 20
    ws.row_dimensions[start_row + 1].height = 26
    _set_border(ws, start_row, start_row + 1, start_col, start_col + 1)
    label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _format_summary_value(value_cell, key, value_format)


@lru_cache(maxsize=256)
def _bookmaker_logo_artifact(codes_key: tuple[str, ...]) -> Optional[tuple[bytes, int, int]]:
    try:
        from PIL import Image as PILImage
    except ImportError:
        return None

    tiles = []
    for code in codes_key[:BOOKMAKER_LOGO_MAX_PER_CELL]:
        filename = BOOKMAKER_LOGO_FILES.get(code)
        if not filename:
            continue
        path = BOOKMAKER_LOGO_DIR / filename
        if not path.is_file():
            continue
        try:
            with PILImage.open(path) as raw_image:
                image = raw_image.convert("RGBA")
                image.thumbnail((BOOKMAKER_LOGO_MAX_WIDTH, BOOKMAKER_LOGO_MAX_HEIGHT), PILImage.LANCZOS)
                tiles.append(image.copy())
        except Exception:
            continue

    if not tiles:
        return None

    width = sum(tile.width for tile in tiles) + BOOKMAKER_LOGO_GAP * (len(tiles) - 1)
    height = max(tile.height for tile in tiles)
    canvas = PILImage.new("RGBA", (width, height), (255, 255, 255, 0))
    x = 0
    for tile in tiles:
        y = (height - tile.height) // 2
        canvas.alpha_composite(tile, (x, y))
        x += tile.width + BOOKMAKER_LOGO_GAP

    buffer = BytesIO()
    canvas.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue(), width, height


def _unique_bookmaker_logo_codes(codes: Iterable[str]) -> tuple[str, ...]:
    unique_codes: list[str] = []
    seen = set()
    for code in codes:
        logo_code = _bookmaker_logo_code(code=code)
        if logo_code and logo_code not in seen:
            seen.add(logo_code)
            unique_codes.append(logo_code)
    return tuple(unique_codes)


def build_bookmaker_logo_png(codes: Iterable[str]) -> Optional[BookmakerLogoPng]:
    codes_key = _unique_bookmaker_logo_codes(codes)[:BOOKMAKER_LOGO_MAX_PER_CELL]
    if not codes_key:
        return None
    artifact = _bookmaker_logo_artifact(codes_key)
    if not artifact:
        return None

    data, width, height = artifact
    return BookmakerLogoPng(data=data, width=width, height=height, codes=codes_key)


def _fit_logo_cell(ws, row: int, col: int, *, width: int, height: int) -> None:
    row_height = max(float(ws.row_dimensions[row].height or 15), height * 0.75 + 5)
    ws.row_dimensions[row].height = row_height
    letter = get_column_letter(col)
    column_width = max(float(ws.column_dimensions[letter].width or 0), min(max(width / 7 + 2, 10), 26))
    ws.column_dimensions[letter].width = column_width


def _add_bookmaker_logos(
    ws,
    row: int,
    col: int,
    codes: Iterable[str],
    *,
    logo_context: Optional[_LogoRenderContext] = None,
) -> bool:
    codes_key = _unique_bookmaker_logo_codes(codes)[:BOOKMAKER_LOGO_MAX_PER_CELL]
    if not codes_key:
        return False
    artifact = _bookmaker_logo_artifact(codes_key)
    if not artifact:
        return False

    data, width, height = artifact
    if logo_context and logo_context.mode == "google_cell":
        logo_context.icon_cells.append(GoogleSheetIconCell(
            sheet=str(ws.title),
            row=row,
            column=col,
            codes=codes_key,
            width=width,
            height=height,
        ))
        _fit_logo_cell(ws, row, col, width=width, height=height)
        return True

    image = WorksheetImage(BytesIO(data))
    image.width = width
    image.height = height
    ws.add_image(image, ws.cell(row=row, column=col).coordinate)

    _fit_logo_cell(ws, row, col, width=width, height=height)
    return True


def _write_monthly_table(
    ws,
    start_row: int,
    items: list[StatsExportItem],
    *,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> int:
    _write_section_title(ws, start_row, 1, 11, "Помесячная сводка", fill=COLOR_INFO_FILL)
    header_row = start_row + 1
    headers = [
        "Месяц",
        "Ставки",
        "Побед",
        "Пораж.",
        "Возврат",
        "Выкуп",
        "Проход",
        "Ср. коэфф.",
        f"Оборот, {value_label}",
        f"Прибыль, {value_label}",
        "ROI",
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=header_row, column=col, value=header)
    _style_range_header(ws, header_row, 1, len(headers), COLOR_HEADER_SAGE)

    monthly_rows = _monthly_rows(items, historical=historical)
    for offset, row_data in enumerate(monthly_rows, 1):
        row = header_row + offset
        values = [
            row_data["label"],
            row_data["bets"],
            row_data["wins"],
            row_data["losses"],
            row_data["refunds"],
            row_data["buyouts"],
            row_data["winrate"],
            row_data["average_coefficient"],
            row_data["turnover"],
            row_data["profit"],
            row_data["roi"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
        for col, key in [(7, "winrate"), (8, "average_coefficient"), (9, "turnover"), (10, "profit"), (11, "roi")]:
            _format_summary_value(ws.cell(row=row, column=col), key, value_format)
        if row_data["label"] == "ИТОГО":
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).font = Font(bold=True)
                ws.cell(row=row, column=col).fill = _solid_fill(COLOR_TOTAL_FILL)
        elif offset % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).fill = _solid_fill(COLOR_SURFACE_ALT)

    end_row = header_row + len(monthly_rows)
    _set_border(ws, header_row, end_row, 1, len(headers))
    return end_row + 3


def _write_compact_breakdowns(
    ws,
    start_row: int,
    items: list[StatsExportItem],
    *,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
    logo_context: Optional[_LogoRenderContext] = None,
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> None:
    bookmaker_rows = _breakdown_rows(items, kind="bookmaker", historical=historical)
    sport_rows = _breakdown_rows(items, kind="sport", historical=historical)

    _write_section_title(ws, start_row, 1, 6, "Букмекеры", fill=COLOR_SAGE_FILL)
    _write_section_title(ws, start_row, 8, 13, "Виды спорта", fill=COLOR_INFO_FILL)
    headers = ["Иконка", "БК", "Ставки", "Проход", f"Прибыль, {value_label}", "ROI"]
    sport_headers = ["Иконка", "Вид спорта", "Ставки", "Проход", f"Прибыль, {value_label}", "ROI"]
    for idx, header in enumerate(headers, 1):
        ws.cell(row=start_row + 1, column=idx, value=header)
    for idx, header in enumerate(sport_headers, 8):
        ws.cell(row=start_row + 1, column=idx, value=header)
    _style_range_header(ws, start_row + 1, 1, 6, COLOR_HEADER_SAGE)
    _style_range_header(ws, start_row + 1, 8, 13, COLOR_HEADER_STEEL)

    for offset, row_data in enumerate(bookmaker_rows, 2):
        row = start_row + offset
        values = [
            _logo_cell_text(row_data["logo_codes"], row_data["icon"]),
            row_data["label"],
            row_data["bets"],
            row_data["winrate"],
            row_data["profit"],
            row_data["roi"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
        _add_bookmaker_logos(ws, row, 1, row_data["logo_codes"], logo_context=logo_context)
        for col, key in [(4, "winrate"), (5, "profit"), (6, "roi")]:
            _format_summary_value(ws.cell(row=row, column=col), key, value_format)
        if offset % 2 == 1:
            for col in range(1, 7):
                ws.cell(row=row, column=col).fill = _solid_fill(COLOR_SURFACE_ALT)

    for offset, row_data in enumerate(sport_rows, 2):
        row = start_row + offset
        values = [row_data["icon"], row_data["label"], row_data["bets"], row_data["winrate"], row_data["profit"], row_data["roi"]]
        for col, value in enumerate(values, 8):
            ws.cell(row=row, column=col, value=value)
        for col, key in [(11, "winrate"), (12, "profit"), (13, "roi")]:
            _format_summary_value(ws.cell(row=row, column=col), key, value_format)
        if offset % 2 == 1:
            for col in range(8, 14):
                ws.cell(row=row, column=col).fill = _solid_fill(COLOR_SURFACE_ALT)

    max_rows = max(len(bookmaker_rows), len(sport_rows)) + start_row + 1
    _set_border(ws, start_row + 1, max_rows, 1, 6)
    _set_border(ws, start_row + 1, max_rows, 8, 13)


def _setup_statistics_sheet(
    wb: Workbook,
    title: str,
    period_label: str,
    items: list[StatsExportItem],
    *,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
    logo_context: Optional[_LogoRenderContext] = None,
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> None:
    ws = wb.active
    ws.title = "Статистика"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_TITLE
    ws["A1"] = title
    _style_title(ws, 13)
    _write_metadata_row(ws, 2, period_label, value_label=value_label)

    summary = summarize_export_items(items, historical=historical)
    _write_kpi(ws, 4, 1, "Всего ставок", summary["bets"], "bets", value_format=value_format)
    _write_kpi(ws, 4, 4, "Проходимость", summary["winrate"], "winrate", value_format=value_format)
    _write_kpi(ws, 4, 7, f"Прибыль, {value_label}", summary["profit"], "profit", value_format=value_format)
    _write_kpi(ws, 4, 10, "ROI", summary["roi"], "roi", value_format=value_format)
    _write_kpi(ws, 7, 1, f"Оборот, {value_label}", summary["turnover"], "turnover", value_format=value_format)
    _write_kpi(ws, 7, 4, "Ср. коэфф.", summary["average_coefficient"], "average_coefficient", value_format=value_format)
    _write_kpi(ws, 7, 7, "Побед", summary["wins"], "wins", value_format=value_format)
    _write_kpi(ws, 7, 10, "Поражений", summary["losses"], "losses", value_format=value_format)

    next_row = _write_monthly_table(ws, 11, items, value_format=value_format, value_label=value_label, historical=historical)
    _write_compact_breakdowns(
        ws,
        next_row,
        items,
        value_format=value_format,
        value_label=value_label,
        logo_context=logo_context,
        historical=historical,
    )
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 14,
        "B": 14,
        "C": 10,
        "D": 14,
        "E": 14,
        "F": 10,
        "G": 16,
        "H": 18,
        "I": 15,
        "J": 15,
        "K": 12,
        "L": 3,
        "M": 15,
    })
    ws.freeze_panes = "A13"


def _write_breakdown_sheet(
    wb: Workbook,
    items: list[StatsExportItem],
    *,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
    logo_context: Optional[_LogoRenderContext] = None,
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> None:
    ws = wb.create_sheet("Свод по БК и спорту")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_HEADER_SAGE
    _write_section_title(ws, 1, 1, 11, "Свод по букмекерам", fill=COLOR_SAGE_FILL)
    _write_section_title(ws, 1, 13, 23, "Свод по видам спорта", fill=COLOR_INFO_FILL)

    headers = ["Иконка", "БК", "Ставки", "Побед", "Пораж.", "Возврат", "Проход", "Ср. коэфф.", f"Оборот, {value_label}", f"Прибыль, {value_label}", "ROI"]
    sport_headers = ["Иконка", "Вид спорта", "Ставки", "Побед", "Пораж.", "Возврат", "Проход", "Ср. коэфф.", f"Оборот, {value_label}", f"Прибыль, {value_label}", "ROI"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=header)
    for col, header in enumerate(sport_headers, 13):
        ws.cell(row=2, column=col, value=header)
    _style_range_header(ws, 2, 1, 11, COLOR_HEADER_SAGE)
    _style_range_header(ws, 2, 13, 23, COLOR_HEADER_STEEL)

    bookmaker_rows = _breakdown_rows(items, kind="bookmaker", historical=historical)
    sport_rows = _breakdown_rows(items, kind="sport", historical=historical)
    for offset, row_data in enumerate(bookmaker_rows, 3):
        _write_breakdown_detail_row(
            ws,
            offset,
            1,
            row_data,
            value_format=value_format,
            add_bookmaker_logo=True,
            logo_context=logo_context,
        )
    for offset, row_data in enumerate(sport_rows, 3):
        _write_breakdown_detail_row(ws, offset, 13, row_data, value_format=value_format)

    max_row = max(len(bookmaker_rows), len(sport_rows)) + 2
    _set_border(ws, 2, max_row, 1, 11)
    _set_border(ws, 2, max_row, 13, 23)
    _set_auto_filter(ws, 2, max_row, 1, 23)
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 13,
        "B": 18,
        "G": 12,
        "H": 12,
        "I": 15,
        "J": 15,
        "K": 12,
        "L": 3,
        "M": 13,
        "N": 18,
        "S": 12,
        "T": 12,
        "U": 15,
        "V": 15,
        "W": 12,
    })
    ws.freeze_panes = "A3"


def _write_breakdown_detail_row(
    ws,
    row: int,
    start_col: int,
    row_data: dict[str, Any],
    *,
    value_format: str = MONEY_FORMAT,
    add_bookmaker_logo: bool = False,
    logo_context: Optional[_LogoRenderContext] = None,
) -> None:
    values = [
        _logo_cell_text(row_data["logo_codes"], row_data["icon"]) if add_bookmaker_logo else row_data["icon"],
        row_data["label"],
        row_data["bets"],
        row_data["wins"],
        row_data["losses"],
        row_data["refunds"],
        row_data["winrate"],
        row_data["average_coefficient"],
        row_data["turnover"],
        row_data["profit"],
        row_data["roi"],
    ]
    for index, value in enumerate(values):
        ws.cell(row=row, column=start_col + index, value=value)
    if add_bookmaker_logo:
        _add_bookmaker_logos(ws, row, start_col, row_data["logo_codes"], logo_context=logo_context)
    for offset, key in [(6, "winrate"), (7, "average_coefficient"), (8, "turnover"), (9, "profit"), (10, "roi")]:
        _format_summary_value(ws.cell(row=row, column=start_col + offset), key, value_format)


def _write_detail_sheet(
    wb: Workbook,
    items: list[StatsExportItem],
    *,
    include_client: bool,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
    logo_context: Optional[_LogoRenderContext] = None,
) -> None:
    ws = wb.create_sheet("Детально")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_HEADER
    headers = [
        "Месяц",
        "День",
        "№",
        "Вид",
        "Вид спорта",
        "Матч",
        "Иконка БК",
        "БК",
        "Коэфф.",
        "Упал до",
        "Ставка",
        "Ставка, флет",
        "Результат",
        f"Оборот, {value_label}",
        f"Прибыль, {value_label}",
    ]
    if value_label != "флеты":
        headers.append("Прибыль, флеты")
    if include_client:
        headers.append("Клиент")
    headers.extend(["Источник", "ID"])
    header_cols = {header: index + 1 for index, header in enumerate(headers)}

    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    _style_range_header(ws, 1, 1, len(headers), COLOR_HEADER)

    month_counts: dict[str, int] = defaultdict(int)
    row_idx = 2
    current_year: Optional[int] = None
    current_month: Optional[str] = None
    current_day: Optional[str] = None
    group_rows: set[int] = set()
    for item in _sorted_items(items):
        resolved_at_msk = as_moscow_datetime(item.resolved_at)
        day_label = resolved_at_msk.strftime("%d.%m.%Y") if resolved_at_msk else ""
        year_label = resolved_at_msk.year if resolved_at_msk else None
        if year_label and year_label != current_year:
            _write_export_group_row(ws, row_idx, len(headers), str(year_label), level=1, fill=COLOR_GROUP_YEAR)
            group_rows.add(row_idx)
            row_idx += 1
            current_year = year_label
            current_month = None
            current_day = None
        if item.period_label != current_month:
            _write_export_group_row(ws, row_idx, len(headers), item.period_label, level=2, fill=COLOR_GROUP_MONTH)
            group_rows.add(row_idx)
            row_idx += 1
            current_month = item.period_label
            current_day = None
        if day_label and day_label != current_day:
            _write_export_group_row(ws, row_idx, len(headers), day_label, level=3, fill=COLOR_GROUP_DAY)
            group_rows.add(row_idx)
            row_idx += 1
            current_day = day_label

        month_counts[item.period_key] += 1
        values = [
            item.period_label,
            day_label,
            month_counts[item.period_key],
            item.sport_icon,
            item.sport_type,
            item.event_name,
            _logo_cell_text(item.bookmaker_logo_codes, item.bookmaker_icon),
            _compact_list_label(item.bookmaker_names),
            item.coefficient,
            item.odds_dropped_to,
            item.stake_text,
            _flat_stake_for_item(item),
            item.result_label,
            item.turnover,
            item.profit,
        ]
        if value_label != "флеты":
            values.append(_flat_profit_for_item(item))
        if include_client:
            values.append(item.client_name)
        values.extend([item.source_type, item.id])
        for col, value in enumerate(values, 1):
            ws.cell(row=row_idx, column=col, value=value)
        _add_bookmaker_logos(
            ws,
            row_idx,
            header_cols["Иконка БК"],
            item.bookmaker_logo_codes,
            logo_context=logo_context,
        )
        ws.cell(row=row_idx, column=header_cols["Коэфф."]).number_format = COEF_FORMAT
        ws.cell(row=row_idx, column=header_cols["Упал до"]).number_format = COEF_FORMAT
        ws.cell(row=row_idx, column=header_cols["Ставка, флет"]).number_format = FLAT_FORMAT
        ws.cell(row=row_idx, column=header_cols[f"Оборот, {value_label}"]).number_format = value_format
        ws.cell(row=row_idx, column=header_cols[f"Прибыль, {value_label}"]).number_format = value_format
        if "Прибыль, флеты" in header_cols:
            ws.cell(row=row_idx, column=header_cols["Прибыль, флеты"]).number_format = FLAT_FORMAT
        fill = _status_fill(item.status)
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).fill = fill
        ws.row_dimensions[row_idx].outlineLevel = 4
        row_idx += 1

    max_row = max(row_idx - 1, 1)
    if max_row > 1:
        _set_border(ws, 1, max_row, 1, len(headers))
    else:
        _set_border(ws, 1, 1, 1, len(headers))
    for row in group_rows:
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
    center_cols = {
        header_cols["День"],
        header_cols["№"],
        header_cols["Вид"],
        header_cols["Иконка БК"],
        header_cols["Результат"],
    }
    wrap_cols = {
        header_cols["Вид спорта"],
        header_cols["Матч"],
        header_cols["БК"],
    }
    if "Клиент" in header_cols:
        wrap_cols.add(header_cols["Клиент"])
    _polish_table_header(ws, 1, 1, len(headers))
    _polish_table_body(
        ws,
        2,
        max_row,
        len(headers),
        center_cols=center_cols,
        wrap_cols=wrap_cols,
        row_height=34,
    )
    for row in group_rows:
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
    ws.sheet_properties.outlinePr.summaryBelow = False
    _set_auto_filter(ws, 1, max_row, 1, len(headers))
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 16,
        "B": 13,
        "C": 8,
        "D": 8,
        "E": 15,
        "F": 42,
        "G": 18,
        "H": 24,
        "I": 11,
        "J": 11,
        "K": 14,
        "L": 13,
        "M": 14,
        "N": 15,
        "O": 15,
        "P": 15,
        "Q": 20,
        "R": 14,
        "S": 36,
    })
    ws.freeze_panes = "A2"


def _write_historical_sheet(
    wb: Workbook,
    historical: Optional[HistoricalStatsSnapshot],
    *,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
) -> None:
    if not historical:
        return

    ws = wb.create_sheet("История")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_HEADER_AMBER
    _write_section_title(ws, 1, 1, 11, "Исторический baseline до 01.07.2026", fill=COLOR_INFO_FILL)

    monthly_headers = [
        "Месяц",
        "Ставки",
        "Побед",
        "Пораж.",
        "Возврат",
        "Проход",
        "Ср. коэфф.",
        f"Оборот, {value_label}",
        f"Прибыль, {value_label}",
        "ROI",
        "Топ БК",
    ]
    for col, header in enumerate(monthly_headers, 1):
        ws.cell(row=3, column=col, value=header)
    _style_range_header(ws, 3, 1, len(monthly_headers), COLOR_HEADER_SAGE)

    for offset, row_data in enumerate(_monthly_rows([], historical=historical), 4):
        values = [
            row_data["label"],
            row_data["bets"],
            row_data["wins"],
            row_data["losses"],
            row_data["refunds"],
            row_data["winrate"],
            row_data["average_coefficient"],
            row_data["turnover"],
            row_data["profit"],
            row_data["roi"],
            next((row.get("top_bookmaker", "") for row in historical.monthly if row.get("period_label") == row_data["label"]), ""),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=offset, column=col, value=value)
        for col, key in [(6, "winrate"), (7, "average_coefficient"), (8, "turnover"), (9, "profit"), (10, "roi")]:
            _format_summary_value(ws.cell(row=offset, column=col), key, value_format)
        if row_data["label"] == "ИТОГО":
            for col in range(1, len(monthly_headers) + 1):
                ws.cell(row=offset, column=col).font = Font(bold=True)
                ws.cell(row=offset, column=col).fill = _solid_fill(COLOR_TOTAL_FILL)

    details_start = 6 + len(historical.monthly)
    _write_section_title(ws, details_start, 1, 12, "Июньские исходные строки из Excel", fill=COLOR_AMBER_FILL)
    detail_headers = [
        "Месяц",
        "№",
        "Вид",
        "Вид спорта",
        "Матч",
        "БК",
        "Коэфф.",
        "Ставка",
        "Результат",
        f"Оборот, {value_label}",
        f"Прибыль, {value_label}",
        "Источник",
    ]
    header_row = details_start + 1
    for col, header in enumerate(detail_headers, 1):
        ws.cell(row=header_row, column=col, value=header)
    _style_range_header(ws, header_row, 1, len(detail_headers), COLOR_HEADER_AMBER)

    for offset, detail in enumerate(historical.details, header_row + 1):
        values = [
            detail.get("period_label", ""),
            detail.get("source_row_number", ""),
            detail.get("sport_icon", ""),
            detail.get("sport_type", ""),
            detail.get("event_name", ""),
            detail.get("bookmaker_name", ""),
            detail.get("coefficient", ""),
            detail.get("outcome", ""),
            _status_label(str(detail.get("status", ""))),
            detail.get("turnover", Decimal("0")),
            detail.get("profit", Decimal("0")),
            detail.get("source_file", ""),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=offset, column=col, value=value)
        ws.cell(row=offset, column=7).number_format = COEF_FORMAT
        ws.cell(row=offset, column=10).number_format = value_format
        ws.cell(row=offset, column=11).number_format = value_format
        fill = _status_fill(str(detail.get("status", "")))
        for col in range(1, len(detail_headers) + 1):
            ws.cell(row=offset, column=col).fill = fill

    max_row = max(header_row + len(historical.details), 3 + len(historical.monthly) + 1)
    _set_border(ws, 3, 3 + len(historical.monthly) + 1, 1, len(monthly_headers))
    _set_border(ws, header_row, max_row, 1, len(detail_headers))
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 16,
        "B": 10,
        "C": 10,
        "D": 16,
        "E": 42,
        "F": 20,
        "G": 12,
        "H": 28,
        "I": 14,
        "J": 15,
        "K": 15,
        "L": 24,
    })
    ws.freeze_panes = "A4"


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        current_width = float(ws.column_dimensions[letter].width or 0)
        ws.column_dimensions[letter].width = max(current_width, min(max(max_length + 2, 10), 48))


def _result_label_short(result: str) -> str:
    if result == "win":
        return "П"
    if result == "loss":
        return "Л"
    return str(result or "")


def _streak_label(streak_type: Optional[str], count: int) -> str:
    if not streak_type or count <= 0:
        return ""
    label = "побед" if streak_type == "win" else "пораж."
    return f"{count} {label}"


def _yes_no(value: bool) -> str:
    return "Да" if value else "Нет"


def _excel_datetime(value: Optional[datetime]) -> Optional[datetime]:
    moscow_value = as_moscow_datetime(value)
    return moscow_value.replace(tzinfo=None) if moscow_value else None


def _client_recent_metric_datetime(row: ClientRecentBetExportRow) -> Optional[datetime]:
    if row.status in {"win", "loss", "refund"}:
        return as_moscow_datetime(row.resolved_at) or as_moscow_datetime(row.taken_at)
    return as_moscow_datetime(row.taken_at)


def _client_recent_period_values(row: ClientRecentBetExportRow) -> tuple[str, str]:
    value = _client_recent_metric_datetime(row)
    if not value:
        return "", ""
    return _month_label(value), value.strftime("%d.%m.%Y")


def _client_recent_flat_stake(row: ClientRecentBetExportRow) -> int:
    return 1 if row.status in {"win", "loss"} else 0


def _client_recent_flat_profit(row: ClientRecentBetExportRow) -> float:
    if row.status == "win":
        return float((Decimal(str(row.coefficient or "0")) - Decimal("1")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    if row.status == "loss":
        return -1.0
    return 0.0


def _split_bookmaker_names(value: str) -> list[str]:
    return [item.strip() for item in str(value or "").split(",") if item.strip()]


def _top_label(values: Iterable[str], default: str = "") -> str:
    counts: dict[str, int] = defaultdict(int)
    labels: dict[str, str] = {}
    for value in values:
        clean = str(value or "").strip()
        if not clean:
            continue
        key = clean.lower()
        counts[key] += 1
        labels.setdefault(key, clean)
    if not counts:
        return default
    key = sorted(counts, key=lambda item: (-counts[item], labels[item].lower()))[0]
    return labels[key]


def _client_period_summary(rows: list[ClientRecentBetExportRow]) -> dict[str, Any]:
    wins = sum(1 for row in rows if row.status == "win")
    losses = sum(1 for row in rows if row.status == "loss")
    refunds = sum(1 for row in rows if row.status == "refund")
    pending = sum(1 for row in rows if row.status == "pending")
    settled = wins + losses
    profit = sum((_client_recent_flat_profit(row) for row in rows), 0.0)
    coefficient_rows = [row for row in rows if row.status in {"win", "loss"} and Decimal(str(row.coefficient or "0")) > 0]
    coefficient_sum = sum((Decimal(str(row.coefficient or "0")) for row in coefficient_rows), Decimal("0"))
    top_sport = _top_label((row.sport_type for row in rows), "Без спорта")
    bookmaker_names = [name for row in rows for name in _split_bookmaker_names(row.bookmaker_names)]
    top_bookmaker = _top_label(bookmaker_names, "Без БК")
    top_bookmaker_logo_codes = []
    for row in rows:
        if top_bookmaker in _split_bookmaker_names(row.bookmaker_names):
            top_bookmaker_logo_codes = list(row.bookmaker_logo_codes)
            break

    return {
        "clients": len({row.user_id for row in rows}),
        "bets": len(rows),
        "wins": wins,
        "losses": losses,
        "refunds": refunds,
        "pending": pending,
        "winrate": (wins / settled) if settled else 0,
        "roi": (profit / settled) if settled else 0,
        "profit": profit,
        "average_coefficient": float((coefficient_sum / Decimal(len(coefficient_rows))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)) if coefficient_rows else 0,
        "sport_icon": SPORT_ICONS.get(top_sport, "✨"),
        "top_sport": top_sport,
        "top_bookmaker": top_bookmaker,
        "top_bookmaker_logo_codes": top_bookmaker_logo_codes,
    }


def _client_period_rows(rows: list[ClientRecentBetExportRow], *, kind: str) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    for row in rows:
        date_value = _client_recent_metric_datetime(row)
        if not date_value:
            continue
        if kind == "month":
            key = date_value.strftime("%Y-%m")
            label = _month_label(date_value)
        else:
            key = date_value.strftime("%Y-%m-%d")
            label = date_value.strftime("%d.%m.%Y")
        groups.setdefault(key, {"label": label, "rows": []})["rows"].append(row)

    period_rows = [
        {
            "label": group["label"],
            **_client_period_summary(group["rows"]),
        }
        for key, group in sorted(groups.items())
    ]
    if period_rows:
        period_rows.append({"label": "ИТОГО", **_client_period_summary(rows)})
    return period_rows


def _client_balance_bucket_value(matches_remaining: int, guarantee_active: bool) -> str:
    if guarantee_active:
        return "Гарантия"
    if matches_remaining < 0:
        return "Долг"
    if matches_remaining == 0:
        return "0 матчей"
    if matches_remaining <= 2:
        return "1-2 матча"
    return "3+ матча"


def _client_balance_bucket(row: ClientInfoExportRow | ClientRecentBetExportRow) -> str:
    return _client_balance_bucket_value(row.matches_remaining, row.guarantee_active)


def _client_recommended_action_value(matches_remaining: int, guarantee_active: bool, situation_tone: str = "") -> str:
    if matches_remaining < 0:
        return "Разобрать долг"
    if guarantee_active:
        return "Закрыть гарантию"
    if matches_remaining == 0:
        return "Продлить абонемент"
    if matches_remaining <= 2:
        return "Допродать матчи"
    if situation_tone == "danger":
        return "Проверить просадку"
    if situation_tone == "warning":
        return "Поддержать клиента"
    return "Наблюдать"


def _client_recommended_action(row: ClientInfoExportRow | ClientRecentBetExportRow) -> str:
    return _client_recommended_action_value(
        row.matches_remaining,
        row.guarantee_active,
        getattr(row, "situation_tone", ""),
    )


def _client_contact_channel_value(vk_user_id: str, username: str, phone: str, is_web_only: bool) -> str:
    if vk_user_id:
        return "VK"
    if username:
        return "Telegram"
    if phone:
        return "Телефон"
    if is_web_only:
        return "Web/VK"
    return "Не указан"


def _client_contact_channel(row: ClientInfoExportRow | ClientRecentBetExportRow) -> str:
    return _client_contact_channel_value(row.vk_user_id, row.username, row.phone, row.is_web_only)


def _client_priority_score(row: ClientInfoExportRow) -> tuple[int, int, float, str]:
    if row.matches_remaining < 0:
        priority = 0
    elif row.guarantee_active:
        priority = 1
    elif row.matches_remaining == 0:
        priority = 2
    elif row.matches_remaining <= 2:
        priority = 3
    elif row.situation_tone in {"danger", "warning"}:
        priority = 4
    else:
        priority = 5
    return (priority, row.matches_remaining, -row.profit_units, row.client_name.lower())


def filter_client_info_export_rows(
    rows: list[ClientInfoExportRow],
    *,
    q: Optional[str],
    activity: str,
    group: Optional[str],
    tag: Optional[str],
    bookmaker_id: Optional[int] = None,
) -> list[ClientInfoExportRow]:
    clean_q = (q or "").strip().lower()
    clean_group = (group or "").strip()
    clean_tag = (tag or "").strip()
    clean_bookmaker_id = bookmaker_id if isinstance(bookmaker_id, int) and bookmaker_id > 0 else None

    def row_matches(row: ClientInfoExportRow) -> bool:
        if clean_q:
            searchable = " ".join([
                str(row.user_id),
                row.client_name,
                row.username,
                row.phone,
                row.vk_user_id,
                row.bookmaker_names,
                row.client_group,
                row.client_tag,
                row.ab_group,
                _client_balance_bucket(row),
                _client_recommended_action(row),
                _client_contact_channel(row),
                str(row.matches_remaining),
                row.situation_label,
                row.situation_description,
            ]).lower()
            if clean_q not in searchable:
                return False
        if activity == "active" and not (row.matches_remaining > 0 or row.guarantee_active):
            return False
        if activity == "empty" and (row.guarantee_active or row.matches_remaining > 0):
            return False
        if activity == "guarantee" and not row.guarantee_active:
            return False
        if clean_group and clean_group != "all" and row.client_group != clean_group:
            return False
        if clean_tag and clean_tag != "all" and row.client_tag != clean_tag:
            return False
        if clean_bookmaker_id and clean_bookmaker_id not in row.bookmaker_ids:
            return False
        return True

    return [row for row in rows if row_matches(row)]


def filter_client_recent_export_rows(
    rows: list[ClientRecentBetExportRow],
    user_ids: set[int],
) -> list[ClientRecentBetExportRow]:
    return [row for row in rows if row.user_id in user_ids]


def _style_chart(chart, *, title: str, height: float = 7.2, width: float = 13.5) -> None:
    chart.title = title
    chart.height = height
    chart.width = width
    chart.style = 10
    chart.legend.position = "r"


def _add_client_overview_charts(
    ws,
    *,
    situation_start_row: int,
    situation_end_row: int,
    balance_start_row: int,
    balance_end_row: int,
    recent_start_row: int,
    recent_end_row: int,
) -> None:
    if situation_end_row >= situation_start_row:
        chart = PieChart()
        labels = Reference(ws, min_col=1, min_row=situation_start_row, max_row=situation_end_row)
        data = Reference(ws, min_col=2, min_row=situation_start_row - 1, max_row=situation_end_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        _style_chart(chart, title="Клиенты по ситуациям", height=7.0, width=12.0)
        ws.add_chart(chart, "P3")

    if balance_end_row >= balance_start_row:
        chart = BarChart()
        labels = Reference(ws, min_col=8, min_row=balance_start_row, max_row=balance_end_row)
        data = Reference(ws, min_col=9, min_row=balance_start_row - 1, max_row=balance_end_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        _style_chart(chart, title="Остатки абона", height=7.0, width=12.0)
        chart.y_axis.title = "Клиентов"
        ws.add_chart(chart, "P18")

    if recent_end_row >= recent_start_row:
        chart = BarChart()
        labels = Reference(ws, min_col=8, min_row=recent_start_row, max_row=recent_end_row)
        data = Reference(ws, min_col=10, min_row=recent_start_row - 1, max_row=recent_end_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        _style_chart(chart, title="Ставки по дням", height=7.0, width=12.0)
        chart.y_axis.title = "Ставок"
        ws.add_chart(chart, "P33")


def _write_client_overview_sheet(
    wb: Workbook,
    rows: list[ClientInfoExportRow],
    recent_rows: list[ClientRecentBetExportRow],
    period_label: str,
) -> None:
    ws = wb.create_sheet("Обзор", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_TITLE
    ws["A1"] = "CRM-ОТЧЕТ SHAMRAI"
    _style_title(ws, 22)
    _write_metadata_row(ws, 2, period_label, value_label="флеты")

    total_clients = len(rows)
    active_clients = sum(1 for row in rows if row.matches_remaining > 0 or row.guarantee_active)
    empty_clients = sum(1 for row in rows if not row.guarantee_active and row.matches_remaining == 0)
    debt_clients = sum(1 for row in rows if row.matches_remaining < 0)
    guarantee_clients = sum(1 for row in rows if row.guarantee_active)
    low_balance_clients = sum(1 for row in rows if not row.guarantee_active and 0 < row.matches_remaining <= 2)
    remaining_matches = sum(row.matches_remaining for row in rows)
    total_taken = sum(row.total_taken_bets for row in rows)
    settled = sum(row.bets for row in rows)
    wins = sum(row.wins for row in rows)
    losses = sum(row.losses for row in rows)
    pending = sum(row.pending_bets for row in rows)
    refunds = sum(row.refund_bets for row in rows)
    profit = sum(row.profit_units for row in rows)
    coefficient_weight = sum(row.average_coefficient * row.bets for row in rows)
    average_coefficient = (coefficient_weight / settled) if settled else 0
    winrate = (wins / settled) if settled else 0
    roi = (profit / settled) if settled else 0

    ws["A3"] = "Коротко по базе: остатки абона, гарантия, клиенты без матчей и ситуация по результатам."
    ws["A3"].font = Font(color=COLOR_TEXT_MUTED, italic=True)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=12)

    _write_kpi(ws, 5, 1, "Клиенты", total_clients, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 5, 4, "Активные", active_clients, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 5, 7, "Матчей осталось", remaining_matches, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 5, 10, "Взял матчей", total_taken, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 8, 1, "Без матчей", empty_clients, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 8, 4, "Гарантия", guarantee_clients, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 8, 7, "Долг", debt_clients, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 8, 10, "Осталось 1-2", low_balance_clients, "bets", value_format=FLAT_FORMAT)
    _write_kpi(ws, 11, 1, "Проход", winrate, "winrate", value_format=FLAT_FORMAT)
    _write_kpi(ws, 11, 4, "ROI", roi, "roi", value_format=FLAT_FORMAT)
    _write_kpi(ws, 11, 7, "Профит, флеты", profit, "profit", value_format=FLAT_FORMAT)
    _write_kpi(ws, 11, 10, "Ср. кф", average_coefficient, "average_coefficient", value_format=FLAT_FORMAT)

    _write_section_title(ws, 15, 1, 6, "Ситуации клиентов", fill=COLOR_SAGE_FILL)
    situation_headers = ["Ситуация", "Клиентов", "Ставки", "Проход", "ROI", "Профит"]
    for col, header in enumerate(situation_headers, 1):
        ws.cell(row=16, column=col, value=header)
    _style_range_header(ws, 16, 1, len(situation_headers), COLOR_HEADER_SAGE)
    grouped: dict[str, list[ClientInfoExportRow]] = defaultdict(list)
    for row in rows:
        grouped[row.situation_label or "Нет статуса"].append(row)
    situation_start_row = 17
    situation_end_row = 16
    for offset, (label, group_rows) in enumerate(sorted(grouped.items(), key=lambda item: (-len(item[1]), item[0])), situation_start_row):
        group_settled = sum(item.bets for item in group_rows)
        group_wins = sum(item.wins for item in group_rows)
        group_profit = sum(item.profit_units for item in group_rows)
        values = [
            label,
            len(group_rows),
            group_settled,
            (group_wins / group_settled) if group_settled else 0,
            (group_profit / group_settled) if group_settled else 0,
            group_profit,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=offset, column=col, value=value)
        ws.cell(row=offset, column=4).number_format = PERCENT_FORMAT
        ws.cell(row=offset, column=5).number_format = PERCENT_FORMAT
        ws.cell(row=offset, column=6).number_format = FLAT_FORMAT
        fill = _solid_fill(COLOR_SURFACE_ALT if offset % 2 else "FFFFFF")
        for col in range(1, len(situation_headers) + 1):
            ws.cell(row=offset, column=col).fill = fill
        situation_end_row = offset

    _write_section_title(ws, 15, 8, 14, "Срез по остаткам абона", fill=COLOR_AMBER_FILL)
    balance_headers = ["Остаток", "Клиентов", "Матчей", "Гарантия", "Средний ROI", "Профит", "Комментарий"]
    for col, header in enumerate(balance_headers, 8):
        ws.cell(row=16, column=col, value=header)
    _style_range_header(ws, 16, 8, 14, COLOR_HEADER_AMBER)
    bucket_order = ["Долг", "0 матчей", "1-2 матча", "3+ матча", "Гарантия"]
    balance_groups: dict[str, list[ClientInfoExportRow]] = defaultdict(list)
    for row in rows:
        balance_groups[_client_balance_bucket(row)].append(row)
    balance_start_row = 17
    balance_end_row = 16
    for offset, label in enumerate([item for item in bucket_order if balance_groups.get(item)], balance_start_row):
        bucket_rows = balance_groups[label]
        bucket_bets = sum(item.bets for item in bucket_rows)
        bucket_profit = sum(item.profit_units for item in bucket_rows)
        values = [
            label,
            len(bucket_rows),
            sum(item.matches_remaining for item in bucket_rows),
            sum(1 for item in bucket_rows if item.guarantee_active),
            (bucket_profit / bucket_bets) if bucket_bets else 0,
            bucket_profit,
            "Связаться" if label in {"Долг", "0 матчей", "1-2 матча", "Гарантия"} else "Рабочая группа",
        ]
        for col, value in enumerate(values, 8):
            ws.cell(row=offset, column=col, value=value)
        ws.cell(row=offset, column=12).number_format = PERCENT_FORMAT
        ws.cell(row=offset, column=13).number_format = FLAT_FORMAT
        fill = _solid_fill(COLOR_WARNING_FILL if offset % 2 else "FFFFFF")
        for col in range(8, 15):
            ws.cell(row=offset, column=col).fill = fill
        balance_end_row = offset

    activity_start_row = max(situation_end_row, balance_end_row, 23) + 3
    _write_section_title(ws, activity_start_row, 1, 7, "Кому срочно написать", fill=COLOR_ROSE_FILL)
    priority_headers = ["Клиент", "Матчей", "Сегмент", "Действие", "Канал", "A/B", "Почему"]
    for col, header in enumerate(priority_headers, 1):
        ws.cell(row=activity_start_row + 1, column=col, value=header)
    _style_range_header(ws, activity_start_row + 1, 1, 7, COLOR_HEADER_ROSE)
    priority_rows = sorted(rows, key=_client_priority_score)[:12]
    for offset, row_data in enumerate(priority_rows, activity_start_row + 2):
        values = [
            row_data.client_name,
            row_data.matches_remaining,
            _client_balance_bucket(row_data),
            _client_recommended_action(row_data),
            _client_contact_channel(row_data),
            row_data.ab_group or "A",
            row_data.situation_description or row_data.situation_label,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=offset, column=col, value=value)
        fill_color = COLOR_DANGER_FILL if row_data.matches_remaining <= 0 or row_data.situation_tone == "danger" else COLOR_WARNING_FILL
        for col in range(1, 8):
            ws.cell(row=offset, column=col).fill = _solid_fill(fill_color)

    _write_section_title(ws, activity_start_row, 8, 14, "Последние ставки: свод", fill=COLOR_INFO_FILL)
    period_rows = _client_period_rows(recent_rows, kind="day")
    recent_headers = ["День", "Клиентов", "Ставки", "Проход", "ROI", "Профит", "Топ спорт"]
    for col, header in enumerate(recent_headers, 8):
        ws.cell(row=activity_start_row + 1, column=col, value=header)
    _style_range_header(ws, activity_start_row + 1, 8, 14, COLOR_HEADER_STEEL)
    recent_start_row = activity_start_row + 2
    recent_end_row = activity_start_row + 1
    for offset, row_data in enumerate(period_rows[-8:] if period_rows else [], recent_start_row):
        values = [
            row_data["label"],
            row_data["clients"],
            row_data["bets"],
            row_data["winrate"],
            row_data["roi"],
            row_data["profit"],
            f'{row_data["sport_icon"]} {row_data["top_sport"]}',
        ]
        for col, value in enumerate(values, 8):
            ws.cell(row=offset, column=col, value=value)
        ws.cell(row=offset, column=11).number_format = PERCENT_FORMAT
        ws.cell(row=offset, column=12).number_format = PERCENT_FORMAT
        ws.cell(row=offset, column=13).number_format = FLAT_FORMAT
        recent_end_row = offset

    max_row = max(ws.max_row, activity_start_row + 1)
    _set_border(ws, 16, max(situation_end_row, 16), 1, 6)
    _set_border(ws, 16, max(balance_end_row, 16), 8, 14)
    _set_border(ws, activity_start_row + 1, max(activity_start_row + 1 + len(priority_rows), activity_start_row + 1), 1, 7)
    _set_border(ws, activity_start_row + 1, max(recent_end_row, activity_start_row + 1), 8, 14)
    _add_client_overview_charts(
        ws,
        situation_start_row=situation_start_row,
        situation_end_row=situation_end_row,
        balance_start_row=balance_start_row,
        balance_end_row=balance_end_row,
        recent_start_row=recent_start_row,
        recent_end_row=recent_end_row,
    )
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 20,
        "B": 13,
        "C": 13,
        "D": 12,
        "E": 12,
        "F": 15,
        "G": 3,
        "H": 15,
        "I": 13,
        "J": 13,
        "K": 12,
        "L": 12,
        "M": 15,
        "N": 22,
        "O": 3,
        "P": 16,
        "Q": 16,
        "R": 16,
        "S": 16,
        "T": 16,
        "U": 16,
        "V": 16,
    })
    ws.freeze_panes = "A16"


def _write_client_period_sheet(
    wb: Workbook,
    rows: list[ClientRecentBetExportRow],
    *,
    title: str,
    kind: str,
    index: int,
    logo_context: Optional[_LogoRenderContext] = None,
) -> None:
    ws = wb.create_sheet(title, index)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_HEADER_SAGE if kind == "month" else COLOR_HEADER_STEEL
    label_header = "Месяц" if kind == "month" else "День"
    headers = [
        label_header,
        "Клиентов",
        "Ставки",
        "Победы",
        "Поражения",
        "Возвраты",
        "Проход",
        "ROI",
        "Профит, флеты",
        "Ср. кф",
        "Ожидают",
        "Вид",
        "Топ спорт",
        "Иконка БК",
        "Топ БК",
    ]
    header_cols = {header: index + 1 for index, header in enumerate(headers)}
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    _style_range_header(ws, 1, 1, len(headers), COLOR_HEADER)

    period_rows = _client_period_rows(rows, kind=kind)
    for row_idx, row_data in enumerate(period_rows, 2):
        values = [
            row_data["label"],
            row_data["clients"],
            row_data["bets"],
            row_data["wins"],
            row_data["losses"],
            row_data["refunds"],
            row_data["winrate"],
            row_data["roi"],
            row_data["profit"],
            row_data["average_coefficient"],
            row_data["pending"],
            row_data["sport_icon"],
            row_data["top_sport"],
            _logo_cell_text(row_data["top_bookmaker_logo_codes"]),
            row_data["top_bookmaker"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row_idx, column=col, value=value)
        _add_bookmaker_logos(
            ws,
            row_idx,
            header_cols["Иконка БК"],
            row_data["top_bookmaker_logo_codes"],
            logo_context=logo_context,
        )
        ws.cell(row=row_idx, column=header_cols["Проход"]).number_format = PERCENT_FORMAT
        ws.cell(row=row_idx, column=header_cols["ROI"]).number_format = PERCENT_FORMAT
        ws.cell(row=row_idx, column=header_cols["Профит, флеты"]).number_format = FLAT_FORMAT
        ws.cell(row=row_idx, column=header_cols["Ср. кф"]).number_format = COEF_FORMAT
        fill_color = COLOR_SUCCESS_FILL if row_data["profit"] > 0 else COLOR_DANGER_FILL if row_data["profit"] < 0 else COLOR_SURFACE
        if row_data["label"] == "ИТОГО":
            fill_color = COLOR_TOTAL_FILL
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).fill = _solid_fill(fill_color)
            if row_data["label"] == "ИТОГО":
                ws.cell(row=row_idx, column=col).font = Font(bold=True, color=COLOR_TEXT)

    max_row = max(ws.max_row, 1)
    _set_border(ws, 1, max_row, 1, len(headers))
    _set_auto_filter(ws, 1, max_row, 1, len(headers))
    _polish_table_header(ws, 1, 1, len(headers))
    if max_row > 1:
        _polish_table_body(
            ws,
            2,
            max_row,
            len(headers),
            center_cols=set(range(2, 15)),
            wrap_cols={header_cols["Топ спорт"], header_cols["Топ БК"]},
            row_height=34,
        )
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 16,
        "B": 13,
        "C": 12,
        "D": 12,
        "E": 12,
        "F": 12,
        "G": 12,
        "H": 12,
        "I": 15,
        "J": 11,
        "K": 12,
        "L": 8,
        "M": 18,
        "N": 18,
        "O": 22,
    })
    ws.freeze_panes = "A2"


def build_client_info_export_workbook_artifact(
    rows: Iterable[ClientInfoExportRow],
    *,
    period_label: str,
    recent_rows: Optional[Iterable[ClientRecentBetExportRow]] = None,
    logo_mode: LogoRenderMode = "floating",
) -> StatsWorkbookArtifact:
    row_list = list(rows)
    recent_row_list = list(recent_rows or [])
    logo_context = _LogoRenderContext(mode=logo_mode)
    wb = Workbook()
    ws = wb.active
    ws.title = "Клиенты"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_HEADER
    headers = [
        "Период",
        "ID",
        "Клиент",
        "Username",
        "Телефон",
        "VK ID",
        "Web/VK клиент",
        "Группа",
        "Тег",
        "A/B",
        "БК клиента",
        "Дата регистрации",
        "Матчей осталось",
        "Сегмент абона",
        "Гарантия",
        "Рекомендованное действие",
        "Канал связи",
        "Взял матчей всего",
        "Рассчитано матчей",
        "Ожидают расчета",
        "Возвраты",
        "Ставки",
        "Победы",
        "Поражения",
        "Проход",
        "ROI",
        "Профит, флеты",
        "Ср. кф",
        "Текущая серия",
        "Макс. побед",
        "Макс. пораж.",
        "Последние исходы",
        "Ситуация",
        "Описание",
    ]
    header_cols = {header: index + 1 for index, header in enumerate(headers)}
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    _style_range_header(ws, 1, 1, len(headers), COLOR_HEADER)

    for row_idx, row_data in enumerate(row_list, 2):
        values = [
            period_label,
            row_data.user_id,
            row_data.client_name,
            f"@{row_data.username}" if row_data.username else "",
            row_data.phone,
            row_data.vk_user_id,
            _yes_no(row_data.is_web_only),
            row_data.client_group,
            row_data.client_tag,
            row_data.ab_group or "A",
            row_data.bookmaker_names,
            _excel_datetime(row_data.created_at),
            row_data.matches_remaining,
            _client_balance_bucket(row_data),
            _yes_no(row_data.guarantee_active),
            _client_recommended_action(row_data),
            _client_contact_channel(row_data),
            row_data.total_taken_bets,
            row_data.settled_bets,
            row_data.pending_bets,
            row_data.refund_bets,
            row_data.bets,
            row_data.wins,
            row_data.losses,
            row_data.winrate / 100,
            row_data.roi / 100,
            row_data.profit_units,
            row_data.average_coefficient,
            _streak_label(row_data.current_streak_type, row_data.current_streak),
            row_data.max_win_streak,
            row_data.max_loss_streak,
            " ".join(_result_label_short(result) for result in row_data.recent_results),
            row_data.situation_label,
            row_data.situation_description,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row_idx, column=col, value=value)
        ws.cell(row=row_idx, column=header_cols["Дата регистрации"]).number_format = DATE_FORMAT
        ws.cell(row=row_idx, column=header_cols["Проход"]).number_format = PERCENT_FORMAT
        ws.cell(row=row_idx, column=header_cols["ROI"]).number_format = PERCENT_FORMAT
        ws.cell(row=row_idx, column=header_cols["Профит, флеты"]).number_format = FLAT_FORMAT
        ws.cell(row=row_idx, column=header_cols["Ср. кф"]).number_format = COEF_FORMAT
        fill = _tone_fill(row_data.situation_tone)
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).fill = fill

    if row_list:
        _set_border(ws, 1, len(row_list) + 1, 1, len(headers))
    else:
        _set_border(ws, 1, 1, 1, len(headers))
    _set_auto_filter(ws, 1, len(row_list) + 1, 1, len(headers))
    center_cols = {
        header_cols["Период"],
        header_cols["ID"],
        header_cols["Телефон"],
        header_cols["VK ID"],
        header_cols["Web/VK клиент"],
        header_cols["Группа"],
        header_cols["Тег"],
        header_cols["A/B"],
        header_cols["Дата регистрации"],
        header_cols["Матчей осталось"],
        header_cols["Сегмент абона"],
        header_cols["Гарантия"],
        header_cols["Рекомендованное действие"],
        header_cols["Канал связи"],
        header_cols["Взял матчей всего"],
        header_cols["Рассчитано матчей"],
        header_cols["Ожидают расчета"],
        header_cols["Возвраты"],
        header_cols["Ставки"],
        header_cols["Победы"],
        header_cols["Поражения"],
        header_cols["Проход"],
        header_cols["ROI"],
        header_cols["Профит, флеты"],
        header_cols["Ср. кф"],
        header_cols["Текущая серия"],
        header_cols["Макс. побед"],
        header_cols["Макс. пораж."],
        header_cols["Последние исходы"],
        header_cols["Ситуация"],
    }
    _polish_table_header(ws, 1, 1, len(headers))
    if row_list:
        _polish_table_body(
            ws,
            2,
            len(row_list) + 1,
            len(headers),
            center_cols=center_cols,
            wrap_cols={
                header_cols["Клиент"],
                header_cols["БК клиента"],
                header_cols["Ситуация"],
                header_cols["Описание"],
            },
            row_height=44,
        )
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 16,
        "B": 11,
        "C": 24,
        "D": 18,
        "E": 17,
        "F": 15,
        "G": 13,
        "H": 14,
        "I": 14,
        "J": 8,
        "K": 52,
        "L": 18,
        "M": 14,
        "N": 16,
        "O": 12,
        "P": 22,
        "Q": 14,
        "R": 16,
        "S": 17,
        "T": 17,
        "U": 12,
        "V": 12,
        "W": 12,
        "X": 12,
        "Y": 12,
        "Z": 12,
        "AA": 15,
        "AB": 11,
        "AC": 16,
        "AD": 12,
        "AE": 12,
        "AF": 18,
        "AG": 20,
        "AH": 44,
    })
    ws.freeze_panes = "A2"
    _write_client_overview_sheet(wb, row_list, recent_row_list, period_label)
    _write_client_period_sheet(
        wb,
        recent_row_list,
        title="По месяцам",
        kind="month",
        index=2,
        logo_context=logo_context,
    )
    _write_client_period_sheet(
        wb,
        recent_row_list,
        title="По дням",
        kind="day",
        index=3,
        logo_context=logo_context,
    )
    _write_client_recent_bets_sheet(wb, recent_row_list, logo_context=logo_context)
    buffer = BytesIO()
    wb.save(buffer)
    return StatsWorkbookArtifact(xlsx=buffer.getvalue(), icon_cells=list(logo_context.icon_cells))


def build_client_info_export_workbook(
    rows: Iterable[ClientInfoExportRow],
    *,
    period_label: str,
    recent_rows: Optional[Iterable[ClientRecentBetExportRow]] = None,
    logo_mode: LogoRenderMode = "floating",
) -> bytes:
    return build_client_info_export_workbook_artifact(
        rows,
        period_label=period_label,
        recent_rows=recent_rows,
        logo_mode=logo_mode,
    ).xlsx


def _write_client_recent_bets_sheet(
    wb: Workbook,
    rows: list[ClientRecentBetExportRow],
    *,
    logo_context: Optional[_LogoRenderContext] = None,
) -> None:
    ws = wb.create_sheet("Последние ставки")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLOR_HEADER_SAGE
    headers = [
        "Месяц",
        "День",
        "ID",
        "Клиент",
        "Username",
        "Телефон",
        "VK ID",
        "Web/VK клиент",
        "Группа",
        "Тег",
        "Матчей осталось",
        "Сегмент абона",
        "Рекомендованное действие",
        "Канал связи",
        "A/B",
        "Гарантия",
        "Дата взятия",
        "Вид",
        "Вид спорта",
        "Матч",
        "Иконка БК",
        "БК",
        "Коэфф.",
        "Ставка, флет",
        "Ставка",
        "Прибыль, флеты",
        "Результат",
        "Источник",
        "Тип доступа",
        "Матч списан",
        "ID ставки",
    ]
    header_cols = {header: index + 1 for index, header in enumerate(headers)}
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    _style_range_header(ws, 1, 1, len(headers), COLOR_HEADER)

    for row_idx, row_data in enumerate(rows, 2):
        month_label, day_label = _client_recent_period_values(row_data)
        values = [
            month_label,
            day_label,
            row_data.user_id,
            row_data.client_name,
            f"@{row_data.username}" if row_data.username else "",
            row_data.phone,
            row_data.vk_user_id,
            _yes_no(row_data.is_web_only),
            row_data.client_group,
            row_data.client_tag,
            row_data.matches_remaining,
            _client_balance_bucket(row_data),
            _client_recommended_action(row_data),
            _client_contact_channel(row_data),
            row_data.ab_group or "A",
            _yes_no(row_data.guarantee_active),
            _excel_datetime(row_data.taken_at),
            SPORT_ICONS.get(row_data.sport_type, "✨"),
            row_data.sport_type,
            _truncate_cell_text(row_data.event_name, limit=80),
            _logo_cell_text(row_data.bookmaker_logo_codes),
            _compact_list_label(row_data.bookmaker_names),
            row_data.coefficient,
            _client_recent_flat_stake(row_data),
            row_data.outcome,
            _client_recent_flat_profit(row_data),
            row_data.result_label,
            row_data.source_type,
            row_data.access_type,
            _yes_no(row_data.match_charged),
            row_data.bet_id,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=row_idx, column=col, value=value)
        _add_bookmaker_logos(
            ws,
            row_idx,
            header_cols["Иконка БК"],
            row_data.bookmaker_logo_codes,
            logo_context=logo_context,
        )
        ws.cell(row=row_idx, column=header_cols["Дата взятия"]).number_format = DATE_FORMAT
        ws.cell(row=row_idx, column=header_cols["Коэфф."]).number_format = COEF_FORMAT
        ws.cell(row=row_idx, column=header_cols["Ставка, флет"]).number_format = FLAT_FORMAT
        ws.cell(row=row_idx, column=header_cols["Прибыль, флеты"]).number_format = FLAT_FORMAT
        fill = _status_fill(row_data.status)
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).fill = fill

    if rows:
        _set_border(ws, 1, len(rows) + 1, 1, len(headers))
    else:
        _set_border(ws, 1, 1, 1, len(headers))
    _set_auto_filter(ws, 1, len(rows) + 1, 1, len(headers))
    center_cols = {
        header_cols["Месяц"],
        header_cols["День"],
        header_cols["ID"],
        header_cols["Телефон"],
        header_cols["VK ID"],
        header_cols["Web/VK клиент"],
        header_cols["Группа"],
        header_cols["Тег"],
        header_cols["Матчей осталось"],
        header_cols["Сегмент абона"],
        header_cols["Рекомендованное действие"],
        header_cols["Канал связи"],
        header_cols["A/B"],
        header_cols["Гарантия"],
        header_cols["Дата взятия"],
        header_cols["Вид"],
        header_cols["Иконка БК"],
        header_cols["Коэфф."],
        header_cols["Ставка, флет"],
        header_cols["Прибыль, флеты"],
        header_cols["Результат"],
        header_cols["Источник"],
        header_cols["Тип доступа"],
        header_cols["Матч списан"],
    }
    _polish_table_header(ws, 1, 1, len(headers))
    if rows:
        _polish_table_body(
            ws,
            2,
            len(rows) + 1,
            len(headers),
            center_cols=center_cols,
            wrap_cols={
                header_cols["Клиент"],
                header_cols["Вид спорта"],
                header_cols["Матч"],
                header_cols["БК"],
                header_cols["Ставка"],
                header_cols["ID ставки"],
            },
            row_height=36,
        )
    _autosize(ws)
    _set_column_widths(ws, {
        "A": 16,
        "B": 13,
        "C": 11,
        "D": 24,
        "E": 18,
        "F": 17,
        "G": 15,
        "H": 13,
        "I": 14,
        "J": 14,
        "K": 14,
        "L": 16,
        "M": 22,
        "N": 14,
        "O": 8,
        "P": 12,
        "Q": 18,
        "R": 8,
        "S": 15,
        "T": 44,
        "U": 18,
        "V": 24,
        "W": 11,
        "X": 13,
        "Y": 14,
        "Z": 15,
        "AA": 14,
        "AB": 15,
        "AC": 13,
        "AD": 14,
        "AE": 36,
    })
    ws.freeze_panes = "A2"


def build_stats_export_workbook_artifact(
    items: Iterable[StatsExportItem],
    *,
    title: str,
    period_label: str,
    include_client: bool = False,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
    logo_mode: LogoRenderMode = "floating",
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> StatsWorkbookArtifact:
    item_list = _sorted_items(items)
    logo_context = _LogoRenderContext(mode=logo_mode)
    wb = Workbook()
    _setup_statistics_sheet(
        wb,
        title,
        period_label,
        item_list,
        value_format=value_format,
        value_label=value_label,
        logo_context=logo_context,
        historical=historical,
    )
    _write_breakdown_sheet(
        wb,
        item_list,
        value_format=value_format,
        value_label=value_label,
        logo_context=logo_context,
        historical=historical,
    )
    _write_detail_sheet(
        wb,
        item_list,
        include_client=include_client,
        value_format=value_format,
        value_label=value_label,
        logo_context=logo_context,
    )
    _write_historical_sheet(
        wb,
        historical,
        value_format=value_format,
        value_label=value_label,
    )
    buffer = BytesIO()
    wb.save(buffer)
    return StatsWorkbookArtifact(xlsx=buffer.getvalue(), icon_cells=list(logo_context.icon_cells))


def build_stats_export_workbook(
    items: Iterable[StatsExportItem],
    *,
    title: str,
    period_label: str,
    include_client: bool = False,
    value_format: str = MONEY_FORMAT,
    value_label: str = "₽",
    logo_mode: LogoRenderMode = "floating",
    historical: Optional[HistoricalStatsSnapshot] = None,
) -> bytes:
    return build_stats_export_workbook_artifact(
        items,
        title=title,
        period_label=period_label,
        include_client=include_client,
        value_format=value_format,
        value_label=value_label,
        logo_mode=logo_mode,
        historical=historical,
    ).xlsx


def stats_export_period_label(period: str) -> str:
    labels = {
        "week": "Текущая неделя",
        "month": "Текущий месяц",
        "quarter": "Текущий квартал",
        "all": "Весь период",
    }
    return labels.get(period, labels["all"])


async def load_shamrai_export_items(db: AsyncSession, period: str) -> list[StatsExportItem]:
    historical = await load_active_historical_stats_snapshot(db, period)
    query = (
        select(Bet)
        .filter(Bet.publication_type == "forecast", Bet.status.in_(list(EXPORT_STATUSES)), Bet.resolved_at.isnot(None))
        .options(selectinload(Bet.bookmaker), selectinload(Bet.bookmakers))
        .order_by(Bet.resolved_at.asc())
    )
    start = period_start(period)
    if start:
        query = query.filter(Bet.resolved_at >= start)
    elif historical:
        query = query.filter(Bet.resolved_at >= historical.cutoff_at)
    result = await db.execute(query)
    return [
        item
        for bet in result.scalars().all()
        if (item := export_item_from_bet(bet)) is not None
    ]


async def load_author_export_items(
    db: AsyncSession,
    *,
    author_id: int,
    period: str,
    source: str = "all",
) -> list[StatsExportItem]:
    query = (
        select(Bet)
        .filter(
            Bet.author_id == author_id,
            Bet.publication_type == "forecast",
            Bet.status.in_(list(EXPORT_STATUSES)),
            Bet.resolved_at.isnot(None),
        )
        .options(selectinload(Bet.bookmaker), selectinload(Bet.bookmakers))
        .order_by(Bet.resolved_at.asc())
    )
    start = period_start(period)
    if start:
        query = query.filter(Bet.resolved_at >= start)
    result = await db.execute(query)
    items = [
        item
        for bet in result.scalars().all()
        if (item := export_item_from_bet(bet)) is not None
    ]
    if source in {"feed", "private", "paid_set"}:
        items = [item for item in items if item.source_type == source]
    return items


async def load_client_export_groups(db: AsyncSession, period: str) -> list[ClientStatsExportGroup]:
    query = (
        select(User, Bet, user_bets.c.access_type, user_bets.c.match_charged, user_bets.c.taken_at)
        .join(user_bets, user_bets.c.user_id == User.telegram_id)
        .join(Bet, Bet.id == user_bets.c.bet_id)
        .filter(
            User.role.notin_(list(STAFF_ROLES)),
            Bet.publication_type == "forecast",
            Bet.status.in_(list(EXPORT_STATUSES)),
            Bet.resolved_at.isnot(None),
        )
        .options(selectinload(Bet.bookmaker), selectinload(Bet.bookmakers))
        .order_by(User.telegram_id.asc(), Bet.resolved_at.asc())
    )
    start = period_start(period)
    if start:
        query = query.filter(Bet.resolved_at >= start)

    grouped: dict[int, ClientStatsExportGroup] = {}
    for user, bet, access_type, match_charged, _taken_at in (await db.execute(query)).all():
        if not is_paid_client_access(access_type, match_charged):
            continue
        client_name = _display_user(user)
        item = export_item_from_bet(
            bet,
            unit_stake=CLIENT_EXPORT_STAKE,
            client_name=client_name,
            access_type=str(access_type or ""),
        )
        if not item:
            continue
        if user.telegram_id not in grouped:
            grouped[user.telegram_id] = ClientStatsExportGroup(
                user_id=user.telegram_id,
                client_name=client_name,
                username=user.username or "",
                items=[],
            )
        grouped[user.telegram_id].items.append(item)
    return list(grouped.values())


async def load_clients_export_items(db: AsyncSession, period: str) -> list[StatsExportItem]:
    groups = await load_client_export_groups(db, period)
    return [item for group in groups for item in group.items]


async def load_client_info_export_rows(db: AsyncSession, period: str) -> list[ClientInfoExportRow]:
    users_result = await db.execute(
        select(User)
        .filter(User.role.notin_(list(STAFF_ROLES)))
        .options(selectinload(User.bookmakers))
        .order_by(User.created_at.desc(), User.telegram_id.desc())
    )
    users = users_result.scalars().all()
    grouped_items: dict[int, list[dict[str, Any]]] = {user.telegram_id: [] for user in users}
    taken_counts: dict[int, dict[str, int]] = {
        user.telegram_id: {
            "total": 0,
            "settled": 0,
            "pending": 0,
            "refund": 0,
        }
        for user in users
    }

    counts_query = (
        select(User.telegram_id, Bet.status, user_bets.c.access_type, user_bets.c.match_charged)
        .join(user_bets, user_bets.c.user_id == User.telegram_id)
        .join(Bet, Bet.id == user_bets.c.bet_id)
        .filter(User.role.notin_(list(STAFF_ROLES)), Bet.publication_type == "forecast")
    )
    for user_id, bet_status, access_type, match_charged in (await db.execute(counts_query)).all():
        counts = taken_counts.setdefault(user_id, {"total": 0, "settled": 0, "pending": 0, "refund": 0})
        counts["total"] += 1
        if bet_status in {"win", "loss"}:
            counts["settled"] += 1
        elif bet_status == "pending":
            counts["pending"] += 1
        elif bet_status == "refund":
            counts["refund"] += 1

    query = (
        select(User, Bet, user_bets.c.access_type, user_bets.c.match_charged, user_bets.c.taken_at)
        .join(user_bets, user_bets.c.user_id == User.telegram_id)
        .join(Bet, Bet.id == user_bets.c.bet_id)
        .filter(
            User.role.notin_(list(STAFF_ROLES)),
            Bet.publication_type == "forecast",
            Bet.status.in_(["win", "loss"]),
            Bet.resolved_at.isnot(None),
        )
        .options(selectinload(Bet.bookmaker), selectinload(Bet.bookmakers))
        .order_by(User.telegram_id.asc(), Bet.resolved_at.desc())
    )
    start = period_start(period)
    if start:
        query = query.filter(Bet.resolved_at >= start)

    for user, bet, access_type, match_charged, taken_at in (await db.execute(query)).all():
        if not is_paid_client_access(access_type, match_charged):
            continue
        item = stat_item_from_bet(
            bet,
            access_type=access_type,
            match_charged=match_charged,
            taken_at=taken_at,
        )
        if item:
            grouped_items.setdefault(user.telegram_id, []).append(item)

    rows: list[ClientInfoExportRow] = []
    for user in users:
        items = grouped_items.get(user.telegram_id, [])
        summary = summarize_items(items)
        situation = client_situation(summary)
        bookmaker_refs = _bookmaker_refs_for_user(user)
        rows.append(ClientInfoExportRow(
            user_id=user.telegram_id,
            client_name=_display_user(user),
            username=user.username or "",
            phone=user.phone or "",
            vk_user_id=user.vk_user_id or "",
            is_web_only=bool(user.is_web_only),
            bookmaker_names=", ".join(bookmaker["name"] for bookmaker in bookmaker_refs),
            client_group=user.client_group or "",
            client_tag=user.client_tag or "",
            created_at=user.created_at,
            matches_remaining=_user_match_balance(user),
            guarantee_active=bool(user.guarantee_active),
            total_taken_bets=int(taken_counts.get(user.telegram_id, {}).get("total", 0)),
            settled_bets=int(taken_counts.get(user.telegram_id, {}).get("settled", 0)),
            pending_bets=int(taken_counts.get(user.telegram_id, {}).get("pending", 0)),
            refund_bets=int(taken_counts.get(user.telegram_id, {}).get("refund", 0)),
            bets=int(summary["bets"]),
            wins=int(summary["wins"]),
            losses=int(summary["losses"]),
            winrate=float(summary["winrate"]),
            roi=float(summary["roi"]),
            profit_units=float(summary["profit_units"]),
            average_coefficient=float(summary["average_coefficient"]),
            current_streak=int(summary["current_streak"]),
            current_streak_type=summary["current_streak_type"],
            max_win_streak=int(summary["max_win_streak"]),
            max_loss_streak=int(summary["max_loss_streak"]),
            recent_results=last_result_codes(items),
            situation_code=situation["code"],
            situation_label=situation["label"],
            situation_tone=situation["tone"],
            situation_description=situation["description"],
            ab_group=str(getattr(user, "ab_group", "") or ""),
            bookmaker_logo_codes=_bookmaker_logo_codes_for_refs(bookmaker_refs),
            bookmaker_ids=[int(bookmaker["id"]) for bookmaker in bookmaker_refs if bookmaker.get("id") is not None],
        ))

    return sorted(rows, key=lambda row: (
        row.situation_tone != "danger",
        -row.matches_remaining,
        -row.profit_units,
        row.client_name.lower(),
    ))


async def load_client_recent_bet_export_rows(
    db: AsyncSession,
    period: str,
    *,
    limit_per_client: Optional[int] = 50,
) -> list[ClientRecentBetExportRow]:
    query = (
        select(User, Bet, user_bets.c.access_type, user_bets.c.match_charged, user_bets.c.taken_at)
        .join(user_bets, user_bets.c.user_id == User.telegram_id)
        .join(Bet, Bet.id == user_bets.c.bet_id)
        .filter(
            User.role.notin_(list(STAFF_ROLES)),
            Bet.publication_type == "forecast",
            Bet.status.in_(["pending", "win", "loss", "refund"]),
        )
        .options(selectinload(User.bookmakers), selectinload(Bet.bookmaker), selectinload(Bet.bookmakers))
        .order_by(User.telegram_id.asc(), user_bets.c.taken_at.desc(), Bet.created_at.desc())
    )
    start = period_start(period)
    if start:
        query = query.filter(user_bets.c.taken_at >= start)

    rows: list[ClientRecentBetExportRow] = []
    per_user_counts: dict[int, int] = defaultdict(int)
    for user, bet, access_type, match_charged, taken_at in (await db.execute(query)).all():
        if limit_per_client is not None and per_user_counts[user.telegram_id] >= limit_per_client:
            continue
        per_user_counts[user.telegram_id] += 1

        bookmakers = _bookmakers_for_bet(bet)
        bookmaker_names = ", ".join(bookmaker["name"] for bookmaker in bookmakers) or "Без БК"
        delivery_mode = str(getattr(bet, "delivery_mode", None) or "feed")
        rows.append(ClientRecentBetExportRow(
            user_id=user.telegram_id,
            client_name=_display_user(user),
            username=user.username or "",
            phone=user.phone or "",
            vk_user_id=user.vk_user_id or "",
            is_web_only=bool(user.is_web_only),
            client_group=user.client_group or "",
            client_tag=user.client_tag or "",
            matches_remaining=_user_match_balance(user),
            guarantee_active=bool(user.guarantee_active),
            taken_at=taken_at,
            event_name=str(getattr(bet, "event_name", "") or ""),
            sport_type=str(getattr(bet, "sport_type", None) or "Без спорта"),
            bookmaker_names=bookmaker_names,
            coefficient=Decimal(str(getattr(bet, "coefficient", None) or "0")),
            outcome=str(getattr(bet, "outcome", None) or ""),
            status=str(getattr(bet, "status", "") or ""),
            result_label=_status_label(str(getattr(bet, "status", "") or "")),
            resolved_at=getattr(bet, "resolved_at", None),
            source_type="feed" if delivery_mode == "feed" else "paid_set" if delivery_mode == "paid_set" else "private",
            access_type=str(access_type or ""),
            match_charged=bool(match_charged),
            bet_id=str(bet.id),
            ab_group=str(getattr(user, "ab_group", "") or ""),
            bookmaker_logo_codes=_bookmaker_logo_codes_for_refs(bookmakers),
        ))
    return rows
