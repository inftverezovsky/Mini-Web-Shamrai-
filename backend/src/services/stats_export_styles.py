from __future__ import annotations

from typing import Any

from openpyxl.styles import PatternFill

# Workbook style tokens and fill helpers used by stats_export.py.
MONEY_FORMAT = '#,##0" ₽"'
FLAT_FORMAT = '0.00" флет"'
PERCENT_FORMAT = "0.0%"
COEF_FORMAT = "0.00"
DATE_FORMAT = "dd.mm.yyyy hh:mm"

XLSX_CALM_PALETTE = {
    "title": "1F2937",
    "header": "293241",
    "header_sage": "52675F",
    "header_steel": "566B80",
    "header_amber": "7C684B",
    "header_rose": "8A5A5A",
    "text": "111827",
    "text_muted": "64748B",
    "border": "D6DAE0",
    "surface": "F7F8FA",
    "surface_alt": "FAFBFC",
    "kpi_label": "E8EDF3",
    "success_fill": "EFF7F1",
    "danger_fill": "F8EEEE",
    "warning_fill": "F8F3E8",
    "info_fill": "EEF4FA",
    "sage_fill": "EEF5F1",
    "amber_fill": "F7F1E8",
    "rose_fill": "F7EDEE",
    "group_year": "E8ECF2",
    "group_month": "EEF4FA",
    "group_day": "EEF5F1",
    "total_fill": "EAF2F7",
}

COLOR_TITLE = XLSX_CALM_PALETTE["title"]
COLOR_HEADER = XLSX_CALM_PALETTE["header"]
COLOR_HEADER_SAGE = XLSX_CALM_PALETTE["header_sage"]
COLOR_HEADER_STEEL = XLSX_CALM_PALETTE["header_steel"]
COLOR_HEADER_AMBER = XLSX_CALM_PALETTE["header_amber"]
COLOR_HEADER_ROSE = XLSX_CALM_PALETTE["header_rose"]
COLOR_TEXT = XLSX_CALM_PALETTE["text"]
COLOR_TEXT_MUTED = XLSX_CALM_PALETTE["text_muted"]
COLOR_BORDER = XLSX_CALM_PALETTE["border"]
COLOR_SURFACE = XLSX_CALM_PALETTE["surface"]
COLOR_SURFACE_ALT = XLSX_CALM_PALETTE["surface_alt"]
COLOR_KPI_LABEL = XLSX_CALM_PALETTE["kpi_label"]
COLOR_SUCCESS_FILL = XLSX_CALM_PALETTE["success_fill"]
COLOR_DANGER_FILL = XLSX_CALM_PALETTE["danger_fill"]
COLOR_WARNING_FILL = XLSX_CALM_PALETTE["warning_fill"]
COLOR_INFO_FILL = XLSX_CALM_PALETTE["info_fill"]
COLOR_SAGE_FILL = XLSX_CALM_PALETTE["sage_fill"]
COLOR_AMBER_FILL = XLSX_CALM_PALETTE["amber_fill"]
COLOR_ROSE_FILL = XLSX_CALM_PALETTE["rose_fill"]
COLOR_GROUP_YEAR = XLSX_CALM_PALETTE["group_year"]
COLOR_GROUP_MONTH = XLSX_CALM_PALETTE["group_month"]
COLOR_GROUP_DAY = XLSX_CALM_PALETTE["group_day"]
COLOR_TOTAL_FILL = XLSX_CALM_PALETTE["total_fill"]


def _solid_fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _status_fill(status: str) -> PatternFill:
    if status == "win":
        return _solid_fill(COLOR_SUCCESS_FILL)
    if status == "loss":
        return _solid_fill(COLOR_DANGER_FILL)
    if status == "pending":
        return _solid_fill(COLOR_WARNING_FILL)
    return _solid_fill(COLOR_SURFACE)


def _tone_fill(tone: str) -> PatternFill:
    if tone == "danger":
        return _solid_fill(COLOR_DANGER_FILL)
    if tone == "warning":
        return _solid_fill(COLOR_WARNING_FILL)
    if tone == "success":
        return _solid_fill(COLOR_SUCCESS_FILL)
    return _solid_fill(COLOR_SURFACE)


def _profit_fill(value: Any) -> PatternFill:
    try:
        numeric_value = float(value)
    except (TypeError, ValueError):
        numeric_value = 0
    if numeric_value > 0:
        return _solid_fill(COLOR_SUCCESS_FILL)
    if numeric_value < 0:
        return _solid_fill(COLOR_DANGER_FILL)
    return _solid_fill(COLOR_SURFACE)
