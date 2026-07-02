from __future__ import annotations

import hashlib
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Optional
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.models.models import (
    HistoricalStatsBreakdown,
    HistoricalStatsDetail,
    HistoricalStatsImportBatch,
    HistoricalStatsMonthly,
)
from src.services.statistics import normalize_period


MOSCOW_TZ = ZoneInfo("Europe/Moscow")
DEFAULT_HISTORICAL_STATS_SOURCE = Path.home() / "Downloads" / "Chrome" / "Статистика зща все время Shamrai.xlsx"
DEFAULT_HISTORICAL_STATS_CUTOFF = datetime(2026, 7, 1, 0, 0, 0, tzinfo=MOSCOW_TZ)
DEFAULT_UNIT_STAKE_RUB = Decimal("10000.00")

MONTHS_RU = {
    "январь": 1,
    "февраль": 2,
    "март": 3,
    "апрель": 4,
    "май": 5,
    "июнь": 6,
    "июль": 7,
    "август": 8,
    "сентябрь": 9,
    "октябрь": 10,
    "ноябрь": 11,
    "декабрь": 12,
}

BOOKMAKER_ALIASES = {
    "fonbet": ("Фонбет", "fonbet"),
    "фонбет": ("Фонбет", "fonbet"),
    "fb": ("Фонбет", "fonbet"),
    "winline": ("Винлайн", "winline"),
    "винлайн": ("Винлайн", "winline"),
    "wl": ("Винлайн", "winline"),
    "лигаставок": ("Лига Ставок", "ligastavok"),
    "лигаставок": ("Лига Ставок", "ligastavok"),
    "лигас": ("Лига Ставок", "ligastavok"),
    "ls": ("Лига Ставок", "ligastavok"),
    "betboom": ("БетБум", "betboom"),
    "бетбум": ("БетБум", "betboom"),
    "bb": ("БетБум", "betboom"),
    "leon": ("Леон", "leon"),
    "леон": ("Леон", "leon"),
    "olimp": ("Олимпбет", "olimpbet"),
    "олимп": ("Олимпбет", "olimpbet"),
    "olimpbet": ("Олимпбет", "olimpbet"),
    "олимпбет": ("Олимпбет", "olimpbet"),
    "марафон": ("Марафонбет", "marathon"),
    "марафонбет": ("Марафонбет", "marathon"),
    "marathon": ("Марафонбет", "marathon"),
    "marathonbet": ("Марафонбет", "marathon"),
    "zenit": ("Зенит", "zenit"),
    "зенит": ("Зенит", "zenit"),
    "pari": ("Пари", "pari"),
    "пари": ("Пари", "pari"),
    "иностр": ("Иностр", "other"),
    "иностран": ("Иностр", "other"),
    "неизвестно": ("Неизвестно", "other"),
}

STATUS_ALIASES = {
    "победа": "win",
    "win": "win",
    "поражение": "loss",
    "loss": "loss",
    "возврат": "refund",
    "refund": "refund",
}


@dataclass(frozen=True)
class HistoricalMonthlyAggregate:
    period_key: str
    period_label: str
    period_start: datetime
    bets: int
    wins: int
    losses: int
    refunds: int
    turnover_rub: Decimal
    profit_rub: Decimal
    average_coefficient: Decimal
    top_sport: str = ""
    top_bookmaker: str = ""


@dataclass(frozen=True)
class HistoricalBreakdownAggregate:
    dimension: str
    icon: str
    label: str
    normalized_key: str
    bookmaker_code: str
    bets: int
    wins: int
    losses: int
    refunds: int
    turnover_rub: Decimal
    profit_rub: Decimal
    average_coefficient: Decimal = Decimal("0")


@dataclass(frozen=True)
class HistoricalDetailAggregate:
    period_key: str
    period_label: str
    source_row_number: int
    sport_icon: str
    sport_type: str
    event_name: str
    bookmaker_icon: str
    bookmaker_name: str
    bookmaker_code: str
    coefficient: Decimal
    outcome: str
    status: str
    turnover_rub: Decimal
    profit_rub: Decimal
    source_file: str = ""


@dataclass
class ParsedHistoricalStats:
    source_path: Path
    source_filename: str
    source_sha256: str
    cutoff_at: datetime
    unit_stake_rub: Decimal
    monthly: list[HistoricalMonthlyAggregate] = field(default_factory=list)
    breakdowns: list[HistoricalBreakdownAggregate] = field(default_factory=list)
    details: list[HistoricalDetailAggregate] = field(default_factory=list)

    @property
    def total_bets(self) -> int:
        return sum(row.bets for row in self.monthly)

    @property
    def total_wins(self) -> int:
        return sum(row.wins for row in self.monthly)

    @property
    def total_losses(self) -> int:
        return sum(row.losses for row in self.monthly)

    @property
    def total_refunds(self) -> int:
        return sum(row.refunds for row in self.monthly)

    @property
    def total_turnover_rub(self) -> Decimal:
        return _money(sum((row.turnover_rub for row in self.monthly), Decimal("0")))

    @property
    def total_profit_rub(self) -> Decimal:
        return _money(sum((row.profit_rub for row in self.monthly), Decimal("0")))


@dataclass(frozen=True)
class HistoricalImportReport:
    status: str
    source_filename: str
    source_sha256: str
    cutoff_at: datetime
    monthly_rows: int
    breakdown_rows: int
    detail_rows: int
    total_bets: int
    total_profit_rub: Decimal


@dataclass(frozen=True)
class HistoricalStatsSnapshot:
    batch_id: Any
    cutoff_at: datetime
    unit_stake_rub: Decimal
    summary: dict[str, Any]
    monthly: list[dict[str, Any]]
    bookmaker_breakdowns: list[dict[str, Any]]
    sport_breakdowns: list[dict[str, Any]]
    details: list[dict[str, Any]]


def _normalise_key(value: str) -> str:
    return "".join(ch for ch in str(value or "").casefold().replace("ё", "е") if ch.isalnum())


def _money(value: Any) -> Decimal:
    return Decimal(str(value or "0")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _decimal(value: Any, places: str = "0.001") -> Decimal:
    return Decimal(str(value or "0")).quantize(Decimal(places), rounding=ROUND_HALF_UP)


def _int(value: Any) -> int:
    if value is None or value == "":
        return 0
    return int(Decimal(str(value)).to_integral_value(rounding=ROUND_HALF_UP))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_bookmaker(value: Any) -> tuple[str, str]:
    raw = str(value or "").strip()
    key = _normalise_key(raw)
    if key in BOOKMAKER_ALIASES:
        return BOOKMAKER_ALIASES[key]
    return raw or "Неизвестно", "other"


def normalize_status(value: Any) -> str:
    key = str(value or "").strip().casefold()
    return STATUS_ALIASES.get(key, "")


def _period_from_label(value: Any, *, default_year: int = 2026) -> tuple[str, str, datetime]:
    label = str(value or "").strip()
    if not label:
        raise ValueError("Missing period label")
    parts = label.split()
    month_name = parts[0].casefold().replace("ё", "е")
    month = MONTHS_RU.get(month_name)
    if not month:
        raise ValueError(f"Unsupported period label: {label}")
    year = int(parts[1]) if len(parts) > 1 and str(parts[1]).isdigit() else (2025 if month >= 9 else default_year)
    period_key = f"{year:04d}-{month:02d}"
    period_label = f"{parts[0]} {year}"
    period_start = datetime(year, month, 1, tzinfo=MOSCOW_TZ)
    return period_key, period_label, period_start


def _iter_non_empty_rows(ws) -> Iterable[tuple[int, tuple[Any, ...]]]:
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if any(value is not None for value in row):
            yield row_index, row


def _parse_monthly(ws) -> list[HistoricalMonthlyAggregate]:
    rows = list(_iter_non_empty_rows(ws))
    header_index = next(
        (index for index, row in rows if str(row[0] or "").strip() == "Период" and str(row[1] or "").strip() == "Ставок"),
        None,
    )
    if header_index is None:
        return []

    monthly = []
    for row_index, row in rows:
        if row_index <= header_index:
            continue
        first = str(row[0] or "").strip()
        if not first or first.startswith("БК база"):
            break
        period_key, period_label, period_start = _period_from_label(first)
        monthly.append(HistoricalMonthlyAggregate(
            period_key=period_key,
            period_label=period_label,
            period_start=period_start,
            bets=_int(row[1]),
            wins=_int(row[2]),
            losses=_int(row[3]),
            refunds=_int(row[4]),
            turnover_rub=_money(row[6]),
            profit_rub=_money(row[7]),
            average_coefficient=_decimal(row[9]),
            top_sport=str(row[10] or "").strip(),
            top_bookmaker=normalize_bookmaker(row[11])[0] if row[11] else "",
        ))
    return monthly


def _parse_breakdown_block(ws, *, marker: str, dimension: str) -> list[HistoricalBreakdownAggregate]:
    rows = list(_iter_non_empty_rows(ws))
    marker_index = next((index for index, row in rows if str(row[0] or "").strip().startswith(marker)), None)
    if marker_index is None:
        return []

    parsed = []
    for row_index, row in rows:
        if row_index <= marker_index + 1:
            continue
        first = str(row[0] or "").strip()
        label = str(row[1] or "").strip()
        if not first or not label or first.startswith("Виды спорта"):
            break
        display_label, bookmaker_code = normalize_bookmaker(label) if dimension == "bookmaker" else (label, "")
        parsed.append(HistoricalBreakdownAggregate(
            dimension=dimension,
            icon=first,
            label=display_label,
            normalized_key=_normalise_key(display_label),
            bookmaker_code=bookmaker_code,
            bets=_int(row[2]),
            wins=_int(row[3]),
            losses=_int(row[4]),
            refunds=_int(row[5]),
            turnover_rub=_money(row[6]),
            profit_rub=_money(row[7]),
            average_coefficient=Decimal("0"),
        ))
    return parsed


def _parse_details(ws) -> list[HistoricalDetailAggregate]:
    details = []
    rows = list(_iter_non_empty_rows(ws))
    if not rows:
        return details
    headers = {str(value or "").strip(): index for index, value in enumerate(rows[0][1])}

    def value(row: tuple[Any, ...], header: str) -> Any:
        index = headers.get(header)
        return row[index] if index is not None and index < len(row) else None

    for row_index, row in rows[1:]:
        period = value(row, "Период")
        event_name = str(value(row, "Матч") or "").strip()
        if not period or not event_name:
            continue
        status = normalize_status(value(row, "Исход"))
        if not status:
            continue
        period_key, period_label, _period_start = _period_from_label(period)
        bookmaker_name, bookmaker_code = normalize_bookmaker(value(row, "БК"))
        details.append(HistoricalDetailAggregate(
            period_key=period_key,
            period_label=period_label,
            source_row_number=_int(value(row, "№")) or row_index,
            sport_icon=str(value(row, "Вид") or "").strip(),
            sport_type=str(value(row, "Вид спорта") or "").strip(),
            event_name=event_name,
            bookmaker_icon=str(value(row, "Иконка БК") or "").strip(),
            bookmaker_name=bookmaker_name,
            bookmaker_code=bookmaker_code,
            coefficient=_decimal(value(row, "Коэф.")),
            outcome=str(value(row, "Ставка") or "").strip(),
            status=status,
            turnover_rub=_money(value(row, "Оборот")),
            profit_rub=_money(value(row, "Прибыль")),
            source_file=str(value(row, "Файл-источник") or "").strip(),
        ))
    return details


def _top_label(values: Iterable[str], fallback: str = "") -> str:
    counts: dict[str, int] = defaultdict(int)
    for value in values:
        clean = str(value or "").strip()
        if clean:
            counts[clean] += 1
    if not counts:
        return fallback
    return sorted(counts.items(), key=lambda item: (-item[1], item[0]))[0][0]


def _monthly_from_details(details: list[HistoricalDetailAggregate]) -> list[HistoricalMonthlyAggregate]:
    grouped: dict[str, list[HistoricalDetailAggregate]] = defaultdict(list)
    for detail in details:
        grouped[detail.period_key].append(detail)

    monthly = []
    for period_key in sorted(grouped):
        rows = grouped[period_key]
        year, month = [int(part) for part in period_key.split("-")]
        first = rows[0]
        wins = sum(1 for row in rows if row.status == "win")
        losses = sum(1 for row in rows if row.status == "loss")
        refunds = sum(1 for row in rows if row.status == "refund")
        coefficient_count = sum(1 for row in rows if row.coefficient > 0)
        coefficient_sum = sum((row.coefficient for row in rows if row.coefficient > 0), Decimal("0"))
        monthly.append(HistoricalMonthlyAggregate(
            period_key=period_key,
            period_label=first.period_label,
            period_start=datetime(year, month, 1, tzinfo=MOSCOW_TZ),
            bets=len(rows),
            wins=wins,
            losses=losses,
            refunds=refunds,
            turnover_rub=_money(sum((row.turnover_rub for row in rows), Decimal("0"))),
            profit_rub=_money(sum((row.profit_rub for row in rows), Decimal("0"))),
            average_coefficient=_decimal(coefficient_sum / Decimal(coefficient_count) if coefficient_count else 0),
            top_sport=_top_label(row.sport_type for row in rows),
            top_bookmaker=_top_label(row.bookmaker_name for row in rows),
        ))
    return monthly


def _breakdowns_from_details(details: list[HistoricalDetailAggregate], *, dimension: str) -> list[HistoricalBreakdownAggregate]:
    grouped: dict[str, list[HistoricalDetailAggregate]] = defaultdict(list)
    for detail in details:
        label = detail.bookmaker_name if dimension == "bookmaker" else detail.sport_type
        grouped[_normalise_key(label)].append(detail)

    rows = []
    for _key, details_group in grouped.items():
        first = details_group[0]
        label = first.bookmaker_name if dimension == "bookmaker" else first.sport_type
        wins = sum(1 for item in details_group if item.status == "win")
        losses = sum(1 for item in details_group if item.status == "loss")
        refunds = sum(1 for item in details_group if item.status == "refund")
        coefficient_count = sum(1 for item in details_group if item.coefficient > 0)
        coefficient_sum = sum((item.coefficient for item in details_group if item.coefficient > 0), Decimal("0"))
        rows.append(HistoricalBreakdownAggregate(
            dimension=dimension,
            icon=first.bookmaker_icon if dimension == "bookmaker" else first.sport_icon,
            label=label,
            normalized_key=_normalise_key(label),
            bookmaker_code=first.bookmaker_code if dimension == "bookmaker" else "",
            bets=len(details_group),
            wins=wins,
            losses=losses,
            refunds=refunds,
            turnover_rub=_money(sum((item.turnover_rub for item in details_group), Decimal("0"))),
            profit_rub=_money(sum((item.profit_rub for item in details_group), Decimal("0"))),
            average_coefficient=_decimal(coefficient_sum / Decimal(coefficient_count) if coefficient_count else 0),
        ))
    return rows


def _combine_breakdowns(rows: Iterable[HistoricalBreakdownAggregate]) -> list[HistoricalBreakdownAggregate]:
    groups: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        key = (row.dimension, row.normalized_key)
        if key not in groups:
            groups[key] = {
                "dimension": row.dimension,
                "icon": row.icon,
                "label": row.label,
                "normalized_key": row.normalized_key,
                "bookmaker_code": row.bookmaker_code,
                "bets": 0,
                "wins": 0,
                "losses": 0,
                "refunds": 0,
                "turnover_rub": Decimal("0"),
                "profit_rub": Decimal("0"),
                "coefficient_sum": Decimal("0"),
                "coefficient_count": 0,
            }
        group = groups[key]
        group["bets"] += row.bets
        group["wins"] += row.wins
        group["losses"] += row.losses
        group["refunds"] += row.refunds
        group["turnover_rub"] += row.turnover_rub
        group["profit_rub"] += row.profit_rub
        if row.average_coefficient > 0 and row.bets > 0:
            group["coefficient_sum"] += row.average_coefficient * Decimal(row.bets)
            group["coefficient_count"] += row.bets

    combined = []
    for group in groups.values():
        coefficient_count = group["coefficient_count"]
        combined.append(HistoricalBreakdownAggregate(
            dimension=group["dimension"],
            icon=group["icon"],
            label=group["label"],
            normalized_key=group["normalized_key"],
            bookmaker_code=group["bookmaker_code"],
            bets=group["bets"],
            wins=group["wins"],
            losses=group["losses"],
            refunds=group["refunds"],
            turnover_rub=_money(group["turnover_rub"]),
            profit_rub=_money(group["profit_rub"]),
            average_coefficient=_decimal(group["coefficient_sum"] / Decimal(coefficient_count) if coefficient_count else 0),
        ))
    return sorted(combined, key=lambda row: (-row.bets, row.label))


def parse_historical_stats_workbook(
    source: Path | str = DEFAULT_HISTORICAL_STATS_SOURCE,
    *,
    cutoff_at: datetime = DEFAULT_HISTORICAL_STATS_CUTOFF,
    unit_stake_rub: Decimal = DEFAULT_UNIT_STAKE_RUB,
) -> ParsedHistoricalStats:
    source_path = Path(source)
    wb = load_workbook(source_path, data_only=True)
    if "История_до_июня" not in wb.sheetnames:
        raise ValueError("Workbook must contain sheet 'История_до_июня'")
    if "Data" not in wb.sheetnames:
        raise ValueError("Workbook must contain sheet 'Data'")

    monthly = _parse_monthly(wb["История_до_июня"])
    details = _parse_details(wb["Data"])
    detail_monthly = _monthly_from_details(details)
    monthly_by_key = {row.period_key: row for row in monthly}
    for row in detail_monthly:
        monthly_by_key[row.period_key] = row

    pre_breakdowns = [
        *_parse_breakdown_block(wb["История_до_июня"], marker="БК база до июня", dimension="bookmaker"),
        *_parse_breakdown_block(wb["История_до_июня"], marker="Виды спорта база до июня", dimension="sport"),
    ]
    detail_breakdowns = [
        *_breakdowns_from_details(details, dimension="bookmaker"),
        *_breakdowns_from_details(details, dimension="sport"),
    ]

    return ParsedHistoricalStats(
        source_path=source_path,
        source_filename=source_path.name,
        source_sha256=_sha256(source_path),
        cutoff_at=cutoff_at,
        unit_stake_rub=_money(unit_stake_rub),
        monthly=sorted(monthly_by_key.values(), key=lambda row: row.period_key),
        breakdowns=_combine_breakdowns([*pre_breakdowns, *detail_breakdowns]),
        details=sorted(details, key=lambda row: (row.period_key, row.source_row_number)),
    )


async def apply_historical_stats_import(
    db: AsyncSession,
    parsed: ParsedHistoricalStats,
    *,
    apply: bool = False,
) -> HistoricalImportReport:
    report = HistoricalImportReport(
        status="dry_run" if not apply else "applied",
        source_filename=parsed.source_filename,
        source_sha256=parsed.source_sha256,
        cutoff_at=parsed.cutoff_at,
        monthly_rows=len(parsed.monthly),
        breakdown_rows=len(parsed.breakdowns),
        detail_rows=len(parsed.details),
        total_bets=parsed.total_bets,
        total_profit_rub=parsed.total_profit_rub,
    )
    if not apply:
        return report

    existing = (
        await db.execute(
            select(HistoricalStatsImportBatch).filter(HistoricalStatsImportBatch.source_sha256 == parsed.source_sha256)
        )
    ).scalars().first()

    other_batches = (
        await db.execute(
            select(HistoricalStatsImportBatch).filter(HistoricalStatsImportBatch.cutoff_at == parsed.cutoff_at)
        )
    ).scalars().all()
    for batch in other_batches:
        if existing is None or batch.id != existing.id:
            batch.is_active = False

    batch = existing or HistoricalStatsImportBatch(source_sha256=parsed.source_sha256)
    batch.source_filename = parsed.source_filename
    batch.cutoff_at = parsed.cutoff_at
    batch.imported_at = datetime.now(timezone.utc)
    batch.unit_stake_rub = parsed.unit_stake_rub
    batch.is_active = True
    batch.total_bets = parsed.total_bets
    batch.total_wins = parsed.total_wins
    batch.total_losses = parsed.total_losses
    batch.total_refunds = parsed.total_refunds
    batch.total_turnover_rub = parsed.total_turnover_rub
    batch.total_profit_rub = parsed.total_profit_rub
    if existing is None:
        db.add(batch)
        await db.flush()

    for model in [HistoricalStatsMonthly, HistoricalStatsBreakdown, HistoricalStatsDetail]:
        await db.execute(delete(model).where(model.batch_id == batch.id))

    db.add_all([
        HistoricalStatsMonthly(
            batch_id=batch.id,
            period_key=row.period_key,
            period_label=row.period_label,
            period_start=row.period_start,
            bets=row.bets,
            wins=row.wins,
            losses=row.losses,
            refunds=row.refunds,
            turnover_rub=row.turnover_rub,
            profit_rub=row.profit_rub,
            average_coefficient=row.average_coefficient,
            top_sport=row.top_sport,
            top_bookmaker=row.top_bookmaker,
        )
        for row in parsed.monthly
    ])
    db.add_all([
        HistoricalStatsBreakdown(
            batch_id=batch.id,
            dimension=row.dimension,
            icon=row.icon,
            label=row.label,
            normalized_key=row.normalized_key,
            bookmaker_code=row.bookmaker_code,
            bets=row.bets,
            wins=row.wins,
            losses=row.losses,
            refunds=row.refunds,
            turnover_rub=row.turnover_rub,
            profit_rub=row.profit_rub,
            average_coefficient=row.average_coefficient,
        )
        for row in parsed.breakdowns
    ])
    db.add_all([
        HistoricalStatsDetail(
            batch_id=batch.id,
            period_key=row.period_key,
            period_label=row.period_label,
            source_row_number=row.source_row_number,
            sport_icon=row.sport_icon,
            sport_type=row.sport_type,
            event_name=row.event_name,
            bookmaker_icon=row.bookmaker_icon,
            bookmaker_name=row.bookmaker_name,
            bookmaker_code=row.bookmaker_code,
            coefficient=row.coefficient,
            outcome=row.outcome,
            status=row.status,
            turnover_rub=row.turnover_rub,
            profit_rub=row.profit_rub,
            source_file=row.source_file,
        )
        for row in parsed.details
    ])
    await db.flush()
    return report


def _snapshot_summary(batch: HistoricalStatsImportBatch) -> dict[str, Any]:
    unit = _money(batch.unit_stake_rub or DEFAULT_UNIT_STAKE_RUB)
    profit_rub = _money(batch.total_profit_rub)
    turnover_rub = _money(batch.total_turnover_rub)
    resolved = int(batch.total_wins or 0) + int(batch.total_losses or 0)
    return {
        "bets": int(batch.total_bets or 0),
        "wins": int(batch.total_wins or 0),
        "losses": int(batch.total_losses or 0),
        "refunds": int(batch.total_refunds or 0),
        "turnover_rub": turnover_rub,
        "profit_rub": profit_rub,
        "profit_units": _money(profit_rub / unit if unit else 0),
        "winrate": (Decimal(batch.total_wins or 0) / Decimal(resolved)) if resolved else Decimal("0"),
        "roi": (profit_rub / turnover_rub) if turnover_rub else Decimal("0"),
    }


async def load_active_historical_stats_snapshot(
    db: AsyncSession,
    period: str,
) -> Optional[HistoricalStatsSnapshot]:
    if normalize_period(period) != "all":
        return None
    if not hasattr(db, "execute"):
        return None
    batch = (
        await db.execute(
            select(HistoricalStatsImportBatch)
            .filter(HistoricalStatsImportBatch.is_active == True)
            .order_by(HistoricalStatsImportBatch.cutoff_at.desc(), HistoricalStatsImportBatch.imported_at.desc())
            .limit(1)
        )
    ).scalars().first()
    if not batch:
        return None

    monthly = (
        await db.execute(
            select(HistoricalStatsMonthly)
            .filter(HistoricalStatsMonthly.batch_id == batch.id)
            .order_by(HistoricalStatsMonthly.period_key.asc())
        )
    ).scalars().all()
    breakdowns = (
        await db.execute(
            select(HistoricalStatsBreakdown)
            .filter(HistoricalStatsBreakdown.batch_id == batch.id)
            .order_by(HistoricalStatsBreakdown.dimension.asc(), HistoricalStatsBreakdown.bets.desc(), HistoricalStatsBreakdown.label.asc())
        )
    ).scalars().all()
    details = (
        await db.execute(
            select(HistoricalStatsDetail)
            .filter(HistoricalStatsDetail.batch_id == batch.id)
            .order_by(HistoricalStatsDetail.period_key.asc(), HistoricalStatsDetail.source_row_number.asc())
        )
    ).scalars().all()

    cutoff_at = batch.cutoff_at
    if cutoff_at and cutoff_at.tzinfo is None:
        cutoff_at = cutoff_at.replace(tzinfo=MOSCOW_TZ)

    return HistoricalStatsSnapshot(
        batch_id=batch.id,
        cutoff_at=cutoff_at,
        unit_stake_rub=_money(batch.unit_stake_rub or DEFAULT_UNIT_STAKE_RUB),
        summary=_snapshot_summary(batch),
        monthly=[
            {
                "period_key": row.period_key,
                "period_label": row.period_label,
                "period_start": row.period_start,
                "bets": int(row.bets or 0),
                "wins": int(row.wins or 0),
                "losses": int(row.losses or 0),
                "refunds": int(row.refunds or 0),
                "turnover": _money(row.turnover_rub),
                "profit": _money(row.profit_rub),
                "average_coefficient": _decimal(row.average_coefficient),
                "top_sport": row.top_sport or "",
                "top_bookmaker": row.top_bookmaker or "",
            }
            for row in monthly
        ],
        bookmaker_breakdowns=[
            _breakdown_snapshot_row(row)
            for row in breakdowns
            if row.dimension == "bookmaker"
        ],
        sport_breakdowns=[
            _breakdown_snapshot_row(row)
            for row in breakdowns
            if row.dimension == "sport"
        ],
        details=[
            {
                "period_key": row.period_key,
                "period_label": row.period_label,
                "source_row_number": int(row.source_row_number or 0),
                "sport_icon": row.sport_icon or "",
                "sport_type": row.sport_type or "",
                "event_name": row.event_name,
                "bookmaker_icon": row.bookmaker_icon or "",
                "bookmaker_name": row.bookmaker_name or "",
                "bookmaker_code": row.bookmaker_code or "",
                "coefficient": _decimal(row.coefficient),
                "outcome": row.outcome or "",
                "status": row.status,
                "turnover": _money(row.turnover_rub),
                "profit": _money(row.profit_rub),
                "source_file": row.source_file or "",
            }
            for row in details
        ],
    )


def _breakdown_snapshot_row(row: HistoricalStatsBreakdown) -> dict[str, Any]:
    return {
        "icon": row.icon or "",
        "label": row.label,
        "normalized_key": row.normalized_key,
        "bookmaker_code": row.bookmaker_code or "",
        "bets": int(row.bets or 0),
        "wins": int(row.wins or 0),
        "losses": int(row.losses or 0),
        "refunds": int(row.refunds or 0),
        "turnover": _money(row.turnover_rub),
        "profit": _money(row.profit_rub),
        "average_coefficient": _decimal(row.average_coefficient),
    }


def _performance_float(value: Any, places: str = "0.01") -> float:
    return float(Decimal(str(value or "0")).quantize(Decimal(places), rounding=ROUND_HALF_UP))


def _performance_summary_from_parts(
    *,
    bets: int,
    wins: int,
    losses: int,
    profit_units: Decimal,
    coefficient_sum: Decimal,
    coefficient_count: int,
    max_win_streak: int = 0,
    max_loss_streak: int = 0,
    current_streak: int = 0,
    current_streak_type: Optional[str] = None,
) -> dict[str, Any]:
    resolved = int(wins or 0) + int(losses or 0)
    return {
        "bets": int(bets or 0),
        "wins": int(wins or 0),
        "losses": int(losses or 0),
        "winrate": round((int(wins or 0) / resolved * 100) if resolved else 0.0, 2),
        "roi": round((float(profit_units) / resolved * 100) if resolved else 0.0, 2),
        "profit_units": _performance_float(profit_units),
        "average_coefficient": _performance_float(
            coefficient_sum / Decimal(coefficient_count)
            if coefficient_count
            else Decimal("0")
        ),
        "max_win_streak": int(max_win_streak or 0),
        "max_loss_streak": int(max_loss_streak or 0),
        "current_streak": int(current_streak or 0),
        "current_streak_type": current_streak_type,
    }


def _performance_summary_parts(summary: dict[str, Any]) -> dict[str, Any]:
    bets = int(summary.get("bets") or 0)
    return {
        "bets": bets,
        "wins": int(summary.get("wins") or 0),
        "losses": int(summary.get("losses") or 0),
        "profit_units": Decimal(str(summary.get("profit_units") or "0")),
        "coefficient_sum": Decimal(str(summary.get("average_coefficient") or "0")) * Decimal(bets),
        "coefficient_count": bets,
    }


def _combine_performance_summaries(
    left: dict[str, Any],
    right: dict[str, Any],
    *,
    streak_source: Optional[dict[str, Any]] = None,
) -> dict[str, Any]:
    left_parts = _performance_summary_parts(left)
    right_parts = _performance_summary_parts(right)
    source = streak_source or right
    return _performance_summary_from_parts(
        bets=left_parts["bets"] + right_parts["bets"],
        wins=left_parts["wins"] + right_parts["wins"],
        losses=left_parts["losses"] + right_parts["losses"],
        profit_units=left_parts["profit_units"] + right_parts["profit_units"],
        coefficient_sum=left_parts["coefficient_sum"] + right_parts["coefficient_sum"],
        coefficient_count=left_parts["coefficient_count"] + right_parts["coefficient_count"],
        max_win_streak=int(source.get("max_win_streak") or 0),
        max_loss_streak=int(source.get("max_loss_streak") or 0),
        current_streak=int(source.get("current_streak") or 0),
        current_streak_type=source.get("current_streak_type"),
    )


def _historical_summary_to_performance(snapshot: HistoricalStatsSnapshot) -> dict[str, Any]:
    summary = snapshot.summary
    coefficient_count = sum(int(row.get("bets") or 0) for row in snapshot.monthly)
    coefficient_sum = sum(
        Decimal(str(row.get("average_coefficient") or "0")) * Decimal(int(row.get("bets") or 0))
        for row in snapshot.monthly
    )
    return _performance_summary_from_parts(
        bets=int(summary.get("bets") or 0),
        wins=int(summary.get("wins") or 0),
        losses=int(summary.get("losses") or 0),
        profit_units=Decimal(str(summary.get("profit_units") or "0")),
        coefficient_sum=coefficient_sum,
        coefficient_count=coefficient_count,
    )


def _historical_row_to_performance_summary(row: dict[str, Any], unit_stake_rub: Decimal) -> dict[str, Any]:
    bets = int(row.get("bets") or 0)
    profit_units = Decimal(str(row.get("profit") or "0")) / unit_stake_rub if unit_stake_rub else Decimal("0")
    coefficient_sum = Decimal(str(row.get("average_coefficient") or "0")) * Decimal(bets)
    return _performance_summary_from_parts(
        bets=bets,
        wins=int(row.get("wins") or 0),
        losses=int(row.get("losses") or 0),
        profit_units=profit_units,
        coefficient_sum=coefficient_sum,
        coefficient_count=bets,
    )


def _historical_month_payload(row: dict[str, Any], unit_stake_rub: Decimal) -> dict[str, Any]:
    key = str(row.get("period_key") or "")
    return {
        "key": key,
        "label": row.get("period_label") or key,
        "summary": _historical_row_to_performance_summary(row, unit_stake_rub),
        "days": [],
    }


def _historical_breakdown_payload(row: dict[str, Any], unit_stake_rub: Decimal) -> dict[str, Any]:
    label = str(row.get("label") or "").strip() or "Без данных"
    key = _normalise_key(label) or str(row.get("normalized_key") or row.get("bookmaker_code") or "none")
    return {
        "key": key,
        "label": label,
        "summary": _historical_row_to_performance_summary(row, unit_stake_rub),
    }


def _merge_performance_months(
    historical_months: list[dict[str, Any]],
    live_months: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    months = {str(month.get("key") or ""): month for month in historical_months if month.get("key")}
    for live_month in live_months:
        key = str(live_month.get("key") or "")
        if not key:
            continue
        existing = months.get(key)
        if not existing:
            months[key] = live_month
            continue
        months[key] = {
            **existing,
            "label": live_month.get("label") or existing.get("label") or key,
            "summary": _combine_performance_summaries(
                existing.get("summary") or {},
                live_month.get("summary") or {},
                streak_source=live_month.get("summary") or {},
            ),
            "days": live_month.get("days") or existing.get("days") or [],
        }
    return [months[key] for key in sorted(months.keys(), reverse=True)]


def _breakdown_merge_key(item: dict[str, Any]) -> str:
    return _normalise_key(str(item.get("label") or item.get("key") or "none")) or "none"


def _merge_performance_breakdown(
    historical_items: list[dict[str, Any]],
    live_items: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    merged: dict[str, dict[str, Any]] = {}
    for item in historical_items:
        merged[_breakdown_merge_key(item)] = item
    for item in live_items:
        key = _breakdown_merge_key(item)
        existing = merged.get(key)
        if not existing:
            merged[key] = item
            continue
        merged[key] = {
            **existing,
            "label": existing.get("label") or item.get("label") or existing.get("key") or key,
            "summary": _combine_performance_summaries(
                existing.get("summary") or {},
                item.get("summary") or {},
                streak_source=item.get("summary") or {},
            ),
        }
    return sorted(
        merged.values(),
        key=lambda item: (
            -int((item.get("summary") or {}).get("bets") or 0),
            str(item.get("label") or ""),
        ),
    )


def merge_historical_performance_payload(
    live_payload: dict[str, Any],
    historical: Optional[HistoricalStatsSnapshot],
) -> dict[str, Any]:
    if not historical:
        return live_payload

    unit_stake_rub = Decimal(str(historical.unit_stake_rub or DEFAULT_UNIT_STAKE_RUB))
    historical_summary = _historical_summary_to_performance(historical)
    live_summary = live_payload.get("summary") or {}
    combined_summary = _combine_performance_summaries(
        historical_summary,
        live_summary,
        streak_source=live_summary,
    )
    live_source_split = live_payload.get("source_split") or {}
    source_split = {
        "all": combined_summary,
        "feed": live_source_split.get("feed") or _performance_summary_from_parts(
            bets=0,
            wins=0,
            losses=0,
            profit_units=Decimal("0"),
            coefficient_sum=Decimal("0"),
            coefficient_count=0,
        ),
        "private": live_source_split.get("private") or _performance_summary_from_parts(
            bets=0,
            wins=0,
            losses=0,
            profit_units=Decimal("0"),
            coefficient_sum=Decimal("0"),
            coefficient_count=0,
        ),
        "paid_set": live_source_split.get("paid_set") or _performance_summary_from_parts(
            bets=0,
            wins=0,
            losses=0,
            profit_units=Decimal("0"),
            coefficient_sum=Decimal("0"),
            coefficient_count=0,
        ),
    }

    return {
        **live_payload,
        "summary": combined_summary,
        "source_split": source_split,
        "timeline": _merge_performance_months(
            [_historical_month_payload(row, unit_stake_rub) for row in historical.monthly],
            list(live_payload.get("timeline") or []),
        ),
        "bookmaker_breakdown": _merge_performance_breakdown(
            [_historical_breakdown_payload(row, unit_stake_rub) for row in historical.bookmaker_breakdowns],
            list(live_payload.get("bookmaker_breakdown") or []),
        ),
        "sport_breakdown": _merge_performance_breakdown(
            [_historical_breakdown_payload(row, unit_stake_rub) for row in historical.sport_breakdowns],
            list(live_payload.get("sport_breakdown") or []),
        ),
    }
